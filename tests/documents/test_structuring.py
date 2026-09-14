"""Structured adapters: Docling JSON, tables, XLSX/CSV and light text."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from open_deep_research.documents import structuring
from open_deep_research.documents.settings import DocumentSettings


def _item(label: str, text: str, *, page: int = 1, name: str | None = None):
    return {
        "self_ref": f"#/{label}s/{name or page}",
        "label": label,
        "prov": [{"page_no": page, "bbox": {"l": 1.0, "t": 2.0, "r": 3.0, "b": 4.0}}],
        "text": {"text": text} if label != "table" else None,
    }


def _table_item(grid: list[list[str]], *, page: int, footnotes: str = "") -> dict:
    return {
        "label": "table",
        "prov": [{"page_no": page, "bbox": {"l": 0, "t": 0, "r": 100, "b": 100}}],
        "table": {
            "data": {"grid": grid, "num_rows": len(grid), "num_cols": len(grid[0])},
            "footnotes": [{"text": footnotes}] if footnotes else [],
        },
    }


def _docling_body(items: list[dict], *, pages: int = 2, status: str = "success") -> dict:
    return {
        "status": status,
        "document": {
            "json_content": {
                "items": items,
                "pages": {str(number): {"size": {"width": 600, "height": 800}}
                          for number in range(1, pages + 1)},
            }
        },
    }


FINANCE_GRID = [
    ["项目", "2025财年", "2024财年"],
    ["营业收入", "48.6亿", "42.1亿"],
    ["毛利率", "31.2%", "29.8%"],
]


def test_from_docling_maps_units_tables_and_merge_candidates():
    body = _docling_body(
        [
            _item("title", "星澜科技 2025 年度财务报表"),
            _item("paragraph", "合并利润表摘要如下。", page=1),
            _table_item(FINANCE_GRID, page=1, footnotes="单位：人民币千元"),
            _item("picture", "chart", page=1),
            _table_item(FINANCE_GRID, page=2, footnotes="单位：人民币千元"),
        ]
    )
    prepared = structuring.from_docling(body)

    types = [unit.unit_type for unit in prepared.units]
    assert types == ["title", "paragraph", "table", "table"]
    title = prepared.units[0]
    assert title.locator["page"] == 1 and title.locator["bbox"] == [1.0, 2.0, 3.0, 4.0]
    table = prepared.units[2]
    assert table.attributes["header"] == FINANCE_GRID[0]
    assert table.attributes["num_rows"] == 3 and table.attributes["num_cols"] == 3
    assert table.attributes["footnotes"] == "单位：人民币千元"
    # Adjacent pages, same column count, repeated header → merge candidate,
    # but both original tables survive untouched as separate units.
    assert any(
        marker.startswith("merge_candidate:")
        for marker in prepared.units[2].attributes["merge_candidates"]
    )
    assert len([u for u in prepared.units if u.unit_type == "table"]) == 2
    assert any(flag.startswith("merge_candidate:") for flag in prepared.quality_flags)
    assert "picture_extracted" in prepared.quality_flags
    assert prepared.page_count == 2


def test_partial_success_is_a_quality_flag_not_a_failure():
    prepared = structuring.from_docling(
        _docling_body([_item("paragraph", "文本", page=1)], status="partial_success")
    )
    assert "docling_partial_success" in prepared.quality_flags


def test_plan_segments_repeats_headers_per_row_group():
    grid = [["指标", "数值"]] + [[f"指标{i}", f"{i}"] for i in range(1, 6)]
    unit = structuring.StructuredUnit(
        unit_type="table",
        locator={"source": "table:1", "page": 1},
        raw_text="\n".join(" | ".join(row) for row in grid),
        index_text="\n".join(" | ".join(row) for row in grid),
        attributes={
            "header": grid[0],
            "num_rows": 6,
            "num_cols": 2,
            "footnotes": "单位：亿元",
        },
    )
    settings = DocumentSettings(table_row_group_size=2)
    prepared = structuring.plan_segments(
        structuring.PreparedDocument(units=[unit]), settings
    )
    assert len(prepared.segment_texts) == 3  # rows 1-2, 3-4, 5
    for text in prepared.segment_texts:
        assert text.startswith("指标 | 数值")  # header repeats in every group
        assert text.endswith("单位：亿元")  # footnotes stay attached
    assert prepared.segment_units == [0, 0, 0]


def test_plan_segments_windows_long_prose():
    unit = structuring.StructuredUnit(
        unit_type="paragraph",
        locator={"source": "paragraph:1"},
        raw_text="段落。" * 2000,
        index_text="段落。" * 2000,
    )
    prepared = structuring.plan_segments(
        structuring.PreparedDocument(units=[unit]), DocumentSettings()
    )
    assert len(prepared.segment_texts) > 1
    assert prepared.segment_units == [0] * len(prepared.segment_texts)


def test_from_xlsx_keeps_formulas_and_missing_cache(tmp_path: Path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "定价"
    sheet.append(["产品", "单价", "数量"])
    sheet.append(["HX-200", 12000, 3])
    sheet.append(["合计", "=B2*C2", None])  # no cached value is ever written
    path = tmp_path / "定价.xlsx"
    workbook.save(path)

    prepared = structuring.from_xlsx(path, DocumentSettings())
    assert len(prepared.units) == 1
    unit = prepared.units[0]
    assert unit.unit_type == "table" and unit.locator["sheet"] == "定价"
    assert unit.attributes["formulas"] == {"B3": "=B2*C2"}
    assert "B3" in unit.attributes["missing_cache"]
    assert any(
        flag.startswith("missing_cached_values:定价") for flag in prepared.quality_flags
    )
    assert "未提供计算结果" in unit.raw_text
    types = unit.attributes["column_types"]
    assert types["单价"] == "integer" and types["产品"] == "text"


def test_from_csv_suggests_column_types(tmp_path: Path):
    path = tmp_path / "rows.csv"
    path.write_text(
        "日期,销量,备注\n2025-01-01,120,首发\n2025-01-02,98,补货\n", encoding="utf-8"
    )
    prepared = structuring.from_csv(path)
    unit = prepared.units[0]
    assert unit.attributes["column_types"]["日期"] == "date"
    assert unit.attributes["column_types"]["销量"] == "integer"
    assert unit.attributes["column_types"]["备注"] == "text"


def test_from_light_text_splits_markdown_headings(tmp_path: Path):
    path = tmp_path / "doc.md"
    path.write_text(
        "# 星澜科技产品手册\n\n概述内容。\n\n## 安装步骤\n\n第一步。\n",
        encoding="utf-8",
    )
    prepared = structuring.from_light_text(path, markdown=True)
    assert prepared.units[0].unit_type == "title"
    headers = [u for u in prepared.units if u.unit_type == "section_header"]
    assert any(u.locator["heading"] == "安装步骤" for u in headers)


def test_mark_merge_candidates_requires_adjacent_pages():
    left = structuring.StructuredUnit(
        unit_type="table", locator={"source": "t1", "page": 1}, raw_text="a | b",
        index_text="a | b", attributes={"header": ["a", "b"], "num_cols": 2, "num_rows": 1},
    )
    far = structuring.StructuredUnit(
        unit_type="table", locator={"source": "t2", "page": 5}, raw_text="a | b",
        index_text="a | b", attributes={"header": ["a", "b"], "num_cols": 2, "num_rows": 1},
    )
    assert structuring.mark_merge_candidates([left, far]) == []
    assert "merge_candidates" not in left.attributes


def test_jsonb_locator_roundtrip_shape():
    unit = structuring.StructuredUnit(
        unit_type="table",
        locator={"source": "table:1:page:2", "page": 2, "bbox": [0.0, 1.0, 2.0, 3.0]},
        raw_text="r", index_text="r",
    )
    assert set(json.loads(json.dumps(unit.locator))) == {"source", "page", "bbox"}


@pytest.mark.parametrize("unknown_label", ["fancy_label", "key_value_region"])
def test_unknown_labels_become_paragraphs(unknown_label):
    body = _docling_body([_item(unknown_label, "内容", page=1)])
    prepared = structuring.from_docling(body)
    assert prepared.units[0].unit_type == "paragraph"


def test_from_docling_core2_separate_arrays():
    body = {
        "status": "success",
        "document": {
            "json_content": {
                "texts": [
                    {"label": "title", "orig": "星澜科技 2025 年度财务报表",
                     "prov": [{"page_no": 1, "bbox": {"l": 1, "t": 9, "r": 2, "b": 8}}]},
                    {"label": "paragraph", "orig": "合并利润表摘要。",
                     "prov": [{"page_no": 1, "bbox": {"l": 1, "t": 7, "r": 2, "b": 6}}]},
                ],
                "tables": [
                    {"data": {"grid": FINANCE_GRID, "num_rows": 3, "num_cols": 3},
                     "prov": [{"page_no": 1, "bbox": {"l": 0, "t": 5, "r": 9, "b": 4}}],
                     "footnotes": [{"text": "单位：人民币千元"}]},
                ],
                "key_value_items": [
                    {"key": {"orig": "审计机构"}, "value": [{"orig": "中衡会计师事务所"}],
                     "prov": [{"page_no": 1, "bbox": {"l": 1, "t": 3, "r": 2, "b": 2}}]},
                ],
                "pictures": [],
                "pages": {"1": {"size": {"width": 600, "height": 800}}},
            }
        },
    }
    prepared = structuring.from_docling(body)
    assert [u.unit_type for u in prepared.units] == [
        "title", "paragraph", "table", "paragraph",
    ]
    kv = prepared.units[3]
    assert kv.raw_text == "审计机构：中衡会计师事务所"
    table = prepared.units[2]
    assert table.attributes["header"] == FINANCE_GRID[0]
    assert table.attributes["footnotes"] == "单位：人民币千元"
    assert prepared.page_count == 1
