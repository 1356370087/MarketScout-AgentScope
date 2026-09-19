"""Run current Office conversion and native parsing with real external components."""

from __future__ import annotations

import asyncio
import json
import os
from pathlib import Path
import subprocess
import uuid


def main():
    for key in list(os.environ):
        if key.lower() in {"http_proxy", "https_proxy", "all_proxy"}:
            os.environ.pop(key)
    root = Path(__file__).resolve().parents[2]
    directory = root / "tmp" / ("m8-office-" + uuid.uuid4().hex[:8])
    directory.mkdir(parents=True)
    from docx import Document
    from pptx import Presentation

    document = Document()
    document.add_heading("Office acceptance", 0)
    document.add_paragraph("Verified document evidence: annual revenue is 120 units.")
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip(
        [*table.rows[0].cells, *table.rows[1].cells], ["Year", "Revenue", "2026", "120"]
    ):
        cell.text = text
    document.save(directory / "sample.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Office acceptance"
    slide.placeholders[
        1
    ].text = "Verified presentation evidence: annual revenue is 120 units."
    slide.notes_slide.notes_text_frame.text = "Acceptance speaker notes"
    deck.save(directory / "sample.pptx")
    # Reuse installed LibreOffice; execute the working-tree converter, not old app code.
    conversion = """
from pathlib import Path
from open_deep_research.documents.office_preview import convert_to_pdf
from open_deep_research.documents.settings import DocumentSettings
for extension in ('docx','pptx'):
    convert_to_pdf(Path('/evidence/sample.'+extension), Path('/evidence/previews/'+extension+'.pdf'), DocumentSettings())
"""
    subprocess.run(
        [
            "docker",
            "run",
            "--rm",
            "--entrypoint",
            "python",
            "-e",
            "PYTHONPATH=/current/src",
            "-v",
            f"{root / 'src'}:/current/src:ro",
            "-v",
            f"{directory}:/evidence",
            "insight_forge-document-worker:latest",
            "-c",
            conversion,
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    from dotenv import dotenv_values

    for key, value in dotenv_values(root / ".env").items():
        if key in {"DOCLING_SERVE_URL", "DOCLING_SERVE_API_KEY"} and value:
            os.environ[key] = value
    from open_deep_research.agentscope_runtime.documents import prepare_document
    from open_deep_research.documents.settings import DocumentSettings

    async def verify():
        results = []
        for extension, media in [
            (
                "docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            (
                "pptx",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            ),
        ]:
            path = directory / ("sample." + extension)
            assert (
                (directory / "previews" / (extension + ".pdf"))
                .read_bytes()
                .startswith(b"%PDF")
            )
            for live in (True, False):
                settings = DocumentSettings(
                    storage_dir=directory,
                    docling_base_url=os.environ["DOCLING_SERVE_URL"] if live else "",
                )
                parsed = await prepare_document(
                    path.read_bytes(), path.name, media, settings, version_id=extension
                )
                assert parsed.segment_texts
                assert "120" in " ".join(parsed.segment_texts)
                assert (
                    "office_preview_used" if live else "docling_not_configured"
                ) in parsed.quality_flags
                if extension == "pptx" and live:
                    assert any(unit.unit_type == "notes" for unit in parsed.units)
                results.append(
                    {
                        "format": extension,
                        "docling": live,
                        "segments": len(parsed.segment_texts),
                        "quality_flags": parsed.quality_flags,
                    }
                )
        (directory / "result.json").write_text(
            json.dumps(results, indent=2), encoding="utf-8"
        )
        print(
            json.dumps(
                {"status": "passed", "cases": len(results), "evidence": str(directory)}
            )
        )

    asyncio.run(verify())


if __name__ == "__main__":
    main()
