"""End-to-end contracts for canonical reports and durable publications."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import threading
import unicodedata

import fitz
import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage
from markdown_it import MarkdownIt
from pptx import Presentation

from open_deep_research import server
from open_deep_research.agents.query_engine import QueryEngine
from open_deep_research.events.publications import PublicationEventStore
from open_deep_research.report.canonical import canonicalize_report
from open_deep_research.report.models import (
    CanonicalReport,
    CanonicalSection,
    CodeBlock,
    InlineRun,
    ListBlock,
    ParagraphBlock,
    PublisherTheme,
    RenderedArtifact,
    TableBlock,
)
from open_deep_research.report.publication_store import (
    PublicationJobStore,
    PublisherSettings,
)
from open_deep_research.report.publisher_worker import PublisherWorker
from open_deep_research.report.publishers import (
    PublicationRenderError,
    render_publication,
    validate_rendered_artifact,
)
from open_deep_research.run_context import RunContextStore
from security.auth import get_current_user
from tests.auth_helpers import research_principal

MARKDOWN = """# Market Report

Executive summary with [Primary](https://example.com/source) and
[Rejected](https://rejected.example/path).

## Findings

- First finding
- Second finding

| Vendor | Score |
| --- | ---: |
| A | 9 |

> Evidence remains incomplete.

```python
print("safe")
```
"""


def _canonical():
    return canonicalize_report(
        MARKDOWN,
        run_id="publication-test",
        report_type="decision_brief",
        locale="en-US",
        sources=[{"title": "Primary", "url": "https://example.com/source"}],
    )


def test_canonical_report_keeps_supported_blocks_and_allowlisted_links() -> None:
    report = _canonical()

    assert report.title == "Market Report"
    assert report.source_markdown_sha256 == hashlib.sha256(
        MARKDOWN.encode()
    ).hexdigest()
    assert report.summary_blocks[0].runs[1].href == "https://example.com/source"
    rejected = next(
        run for run in report.summary_blocks[0].runs if "Rejected" in run.text
    )
    assert rejected.href is None
    assert any(isinstance(block, ListBlock) for block in report.sections[0].blocks)
    assert any(isinstance(block, TableBlock) for block in report.sections[0].blocks)


@pytest.mark.asyncio
async def test_file_output_formats_keep_report_generation_binary_free(monkeypatch) -> None:
    from open_deep_research.report import assembly as assembly_module
    from open_deep_research.report import build_report

    async def fake_invoke(model, messages, config, *, span_name, agent_role=None, model_name=None, **_kwargs):
        return AIMessage(content="# Report\n\nBody")

    monkeypatch.setattr(
        assembly_module,
        "invoke_model_with_retry_observability",
        fake_invoke,
    )
    update = await build_report(
        {"messages": [], "research_brief": "brief", "notes": ["finding"]},
        {"configurable": {"output_format": "pdf", "web_pipeline_mode": "legacy", "quality_evaluation_enabled": False}, "metadata": {"run_id": "binary-free"}},
    )

    assert "report_artifacts" not in update
    assert update["canonical_report"]["source_markdown_sha256"] == hashlib.sha256(
        b"# Report\n\nBody"
    ).hexdigest()


@pytest.mark.parametrize("publication_format", ["markdown", "json"])
def test_text_publishers_return_utf8(publication_format: str) -> None:
    artifact = render_publication(
        _canonical(),
        publication_format,
        PublisherTheme(locale="en-US"),
        source_markdown=MARKDOWN,
    )

    assert artifact.content.decode("utf-8")


def test_markdown_fallback_contains_code_fences_and_escapes_active_markup() -> None:
    code_text = "safe\n````\n<script>alert(1)</script>"
    report = CanonicalReport(
        run_id="markdown-fallback",
        title="Report <unsafe> *title*",
        summary_blocks=[
            ParagraphBlock(
                runs=[
                    InlineRun(text="[bad](https://evil.example) <script> * _ `"),
                    InlineRun(text="a``b", code=True),
                ]
            ),
            CodeBlock(text=code_text, language="py`thon<script>"),
        ],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(report, "markdown", PublisherTheme())
    markdown = artifact.content.decode("utf-8")
    tokens = MarkdownIt("commonmark", {"html": True}).parse(markdown)
    fence = next(token for token in tokens if token.type == "fence")
    inline_children = [
        child
        for token in tokens
        for child in (token.children or [])
    ]

    assert fence.markup == "`" * 5
    assert fence.info == "pythonscript"
    assert fence.content == code_text + "\n"
    assert not any(child.type.startswith("html_") for child in inline_children)
    assert not any(child.type == "link_open" for child in inline_children)
    assert next(child for child in inline_children if child.type == "code_inline").content == "a``b"


def test_pdf_docx_and_pptx_publishers_are_valid_packages() -> None:
    report = _canonical()
    theme = PublisherTheme(locale="en-US", font_family="sans")

    pdf = render_publication(report, "pdf", theme)
    pdf_document = fitz.open(stream=pdf.content, filetype="pdf")
    assert pdf_document.page_count >= 1
    assert "Market Report" in "".join(page.get_text() for page in pdf_document)
    pdf_document.close()

    one_pager = render_publication(report, "one_pager", theme)
    one_page_document = fitz.open(stream=one_pager.content, filetype="pdf")
    assert one_page_document.page_count == 1
    one_page_document.close()

    docx = render_publication(report, "docx", theme)
    word = Document(io.BytesIO(docx.content))
    assert word.core_properties.title == "Market Report"
    assert word.tables[0].cell(1, 0).text == "A"

    pptx = render_publication(report, "slides", theme)
    deck = Presentation(io.BytesIO(pptx.content))
    assert len(deck.slides) == pptx.slide_count
    assert deck.slides[0].shapes.title.text == "Market Report"
    assert any(shape.has_table for slide in deck.slides for shape in slide.shapes)
    assert pptx.preview and pptx.preview["slides"]


def test_pptx_wraps_and_paginates_one_extreme_bullet() -> None:
    bullet = "https://example.com/" + "x" * 4_076
    report = CanonicalReport(
        run_id="long-pptx-bullet",
        title="Long bullet",
        summary_blocks=[ParagraphBlock(runs=[InlineRun(text=bullet)])],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(
        report,
        "pptx",
        PublisherTheme(locale="en-US", font_family="sans"),
    )
    deck = Presentation(io.BytesIO(artifact.content))
    preview_bullets = [
        item
        for slide in artifact.preview["slides"]
        for item in slide["bullets"]
    ]

    assert artifact.slide_count and artifact.slide_count > 2
    assert "".join(item.replace("\n", "") for item in preview_bullets) == bullet
    assert max(
        len(line)
        for slide in list(deck.slides)[1:]
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for line in paragraph.text.splitlines()
    ) <= 72


def test_pptx_wraps_cjk_text_by_display_width() -> None:
    bullet = "动力电池回收与再利用产业" * 40
    report = CanonicalReport(
        run_id="cjk-pptx-bullet",
        title="CJK bullet",
        summary_blocks=[ParagraphBlock(runs=[InlineRun(text=bullet)])],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(
        report,
        "pptx",
        PublisherTheme(locale="zh-CN", font_family="cjk_sans"),
    )
    preview_lines = [
        line
        for slide in artifact.preview["slides"][1:]
        for item in slide["bullets"]
        for line in item.splitlines()
    ]

    assert max(
        sum(
            2 if unicodedata.east_asian_width(character) in {"F", "W"} else 1
            for character in line
        )
        for line in preview_lines
    ) <= 72


def test_default_pptx_limit_accepts_standard_long_research_deck(
    monkeypatch,
) -> None:
    monkeypatch.delenv("PUBLISHER_MAX_PPTX_SLIDES", raising=False)
    monkeypatch.delenv("PUBLISHER_MAX_PPTX_PAGES", raising=False)
    report = CanonicalReport(
        run_id="long-research-deck",
        title="Long research deck",
        sections=[
            CanonicalSection(
                id=f"section-{index}",
                title=f"Section {index}",
                blocks=[ParagraphBlock(runs=[InlineRun(text="Finding")])],
            )
            for index in range(48)
        ],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(report, "pptx", PublisherTheme(locale="en-US"))
    settings = PublisherSettings()

    assert artifact.slide_count == 49
    assert artifact.slide_count <= settings.max_pptx_slides
    validate_rendered_artifact(
        report,
        "pptx",
        artifact,
        max_pptx_slides=settings.max_pptx_slides,
    )
    with pytest.raises(PublicationRenderError, match="pptx_slide_limit_exceeded"):
        validate_rendered_artifact(
            report,
            "pptx",
            artifact,
            max_pptx_slides=48,
        )


def test_one_pager_fails_instead_of_rendering_unreadable_text() -> None:
    markdown = "# Long\n\n" + "word " * 50_000
    report = canonicalize_report(markdown, run_id="long")

    with pytest.raises(PublicationRenderError, match="one_page_overflow"):
        render_publication(report, "one_pager", PublisherTheme())


@pytest.mark.parametrize("font_family", ["sans", "serif"])
def test_pdf_wraps_long_title_and_section_heading_without_truncation(
    font_family: str,
) -> None:
    report = CanonicalReport(
        run_id="long-heading",
        title="T" * 500,
        sections=[CanonicalSection(id="section-1", title="S" * 500)],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(
        report,
        "pdf",
        PublisherTheme(locale="en-US", font_family=font_family),
    )
    document = fitz.open(stream=artifact.content, filetype="pdf")
    try:
        text = "".join(page.get_text() for page in document)
    finally:
        document.close()

    assert text.count("T") == 500
    assert text.count("S") == 500


@pytest.mark.parametrize("font_family", ["sans", "serif"])
def test_pdf_wraps_long_section_heading_at_page_start(font_family: str) -> None:
    report = CanonicalReport(
        run_id="long-section-heading",
        title="Report",
        sections=[CanonicalSection(id="section-1", title="S" * 500)],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(
        report,
        "pdf",
        PublisherTheme(locale="en-US", font_family=font_family),
    )
    document = fitz.open(stream=artifact.content, filetype="pdf")
    try:
        text = "".join(page.get_text() for page in document)
    finally:
        document.close()

    assert text.count("S") == 500


@pytest.mark.parametrize("font_family", ["sans", "serif"])
def test_pdf_wraps_long_spaced_headings_without_truncation(font_family: str) -> None:
    title = "qz " * 120
    section_title = "yx " * 90
    report = CanonicalReport(
        run_id="long-spaced-heading",
        title=title,
        sections=[CanonicalSection(id="section-1", title=section_title)],
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(
        report,
        "pdf",
        PublisherTheme(locale="en-US", font_family=font_family),
    )
    document = fitz.open(stream=artifact.content, filetype="pdf")
    try:
        text = "".join(page.get_text() for page in document)
    finally:
        document.close()

    assert text.count("q") == title.count("q")
    assert text.count("y") == section_title.count("y")


def test_docx_bounds_core_title_without_truncating_body_title() -> None:
    title = "T" * 500
    report = CanonicalReport(
        run_id="long-docx-title",
        title=title,
        source_markdown_sha256="0" * 64,
    )

    artifact = render_publication(
        report,
        "docx",
        PublisherTheme(locale="en-US", font_family="sans"),
    )
    word = Document(io.BytesIO(artifact.content))

    assert word.core_properties.title == title[:255]
    assert word.paragraphs[0].text == title


def _completed_run(tmp_path, run_id: str = "publish-api") -> RunContextStore:
    context = RunContextStore(run_id, runs_dir=str(tmp_path))
    context.initialize(
        "user-1",
        {
            "configurable": {"runs_dir": str(tmp_path), "output_format": "pdf"},
            "metadata": {"run_id": run_id, "owner": "user-1"},
        },
    )
    context.write_text_atomic("final_report.md", MARKDOWN)
    context._update_manifest(  # noqa: SLF001 - terminal API fixture
        status="completed",
        title="Market Report",
        result={"status": "success"},
        publication_theme=PublisherTheme(locale="en-US").model_dump(mode="json"),
    )
    return context


def test_job_store_is_idempotent_and_worker_commits_hash_verified_file(
    tmp_path,
) -> None:
    context = _completed_run(tmp_path, "publish-worker")
    store = PublicationJobStore("publish-worker", runs_dir=tmp_path)
    report = context.brief_path.parent.joinpath("final_report.md").read_text()
    digest = hashlib.sha256(report.encode()).hexdigest()
    theme = PublisherTheme(locale="en-US")

    first, created = store.enqueue(
        report_sha256=digest,
        publication_format="pdf",
        theme=theme,
        max_attempts=3,
    )
    second, duplicate_created = store.enqueue(
        report_sha256=digest,
        publication_format="pdf",
        theme=theme,
        max_attempts=3,
    )

    assert created is True
    assert duplicate_created is False
    assert first.publication_id == second.publication_id
    settings = PublisherSettings(runs_dir=tmp_path, max_concurrent_jobs=1)
    assert asyncio.run(
        PublisherWorker(settings, worker_id="publisher-test").run_once()
    ) == 1
    completed = store.get(first.publication_id)
    assert completed is not None and completed.status == "completed"
    assert store.artifact_path(completed).read_bytes().startswith(b"%PDF")
    assert store.canonical_report_path.is_file()


def test_query_engine_persists_markdown_and_canonical_bundle(tmp_path) -> None:
    run_id = "canonical-bundle"
    engine = QueryEngine(
        {
            "configurable": {"runs_dir": str(tmp_path)},
            "metadata": {"run_id": run_id, "owner": "user-1"},
        }
    )
    assert engine.context_store is not None
    engine.context_store.initialize("user-1", engine.config)
    state = {
        "final_report": MARKDOWN,
        "research_brief": "Market Report",
        "sources": [{"title": "Primary", "url": "https://example.com/source"}],
    }

    asyncio.run(engine._write_report_bundle(state))  # noqa: SLF001

    assert (engine.context_store.context_dir / "final_report.md").read_text() == MARKDOWN
    canonical = (engine.context_store.context_dir / "canonical_report.json").read_text()
    assert '"title": "Market Report"' in canonical


@pytest.mark.asyncio
async def test_run_snapshot_offloads_publication_listing(tmp_path, monkeypatch) -> None:
    caller_thread = threading.get_ident()
    worker_thread = caller_thread

    def record_thread(_self):
        nonlocal worker_thread
        worker_thread = threading.get_ident()
        return []

    monkeypatch.setattr(PublicationJobStore, "list", record_thread)

    assert await server._run_publications("offloaded-list", str(tmp_path)) == []  # noqa: SLF001
    assert worker_thread != caller_thread


def test_publication_events_are_independent_and_drop_unknown_payload_fields(
    tmp_path,
) -> None:
    store = PublicationEventStore("event-run", runs_dir=tmp_path)
    first = store.append(
        "publication.queued",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "queued", "local_path": "secret"},
        dedupe_key="pub-1:queued",
    )
    duplicate = store.append(
        "publication.queued",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "different"},
        dedupe_key="pub-1:queued",
    )

    assert first.sequence == duplicate.sequence == 1
    assert "local_path" not in first.payload
    assert store.last_sequence() == 1


def test_publication_event_append_reads_only_log_tail(tmp_path, monkeypatch) -> None:
    store = PublicationEventStore("event-tail-read", runs_dir=tmp_path)
    store.append(
        "publication.queued",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "queued", "attempt": 0},
        dedupe_key="pub-1:queued",
    )

    def fail_full_read(*, repair_tail=True):
        raise AssertionError(f"unexpected full log read: {repair_tail}")

    monkeypatch.setattr(store, "_read_unlocked", fail_full_read)
    second = store.append(
        "publication.started",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "running", "attempt": 1},
        dedupe_key="pub-1:started:1",
    )
    duplicate = store.append(
        "publication.started",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "running", "attempt": 1},
        dedupe_key="pub-1:started:1",
    )

    assert second.sequence == duplicate.sequence == 2


def test_publication_events_sanitize_invalid_payload_types(tmp_path) -> None:
    store = PublicationEventStore("event-types", runs_dir=tmp_path)

    event = store.append(
        "publication.completed",
        publication_id="pub-1",
        payload={
            "format": 7,
            "status": True,
            "attempt": "one",
            "filename": 8,
            "size_bytes": False,
        },
        dedupe_key="pub-1:completed",
    )

    assert event.payload["format"] is None
    assert event.payload["status"] is None
    assert event.payload["attempt"] is None
    assert event.payload["filename"] is None
    assert event.payload["size_bytes"] is None
    assert store.read() == [event]


def test_publication_events_reject_unknown_schema_version(tmp_path) -> None:
    store = PublicationEventStore("event-schema", runs_dir=tmp_path)
    store.append(
        "publication.queued",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "queued", "attempt": 0},
        dedupe_key="pub-1:queued",
    )
    payload = json.loads(store.path.read_text(encoding="utf-8"))
    payload["schema_version"] = 2
    store.path.write_text(json.dumps(payload) + "\n", encoding="utf-8")

    with pytest.raises(ValueError, match="publication_event_log_corrupted"):
        store.read()


def test_publication_event_tail_repair_is_atomic(tmp_path, monkeypatch) -> None:
    from open_deep_research.events import publications as events_module

    store = PublicationEventStore("event-atomic-repair", runs_dir=tmp_path)
    store.append(
        "publication.queued",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "queued", "attempt": 0},
        dedupe_key="pub-1:queued",
    )
    with store.path.open("ab") as handle:
        handle.write(b'{"partial"')
    original = store.path.read_bytes()

    def fail_replace(_source, _target):
        raise OSError("simulated crash before replace")

    monkeypatch.setattr(events_module.os, "replace", fail_replace)
    with pytest.raises(OSError, match="simulated crash"):
        store.read()
    assert store.path.read_bytes() == original


@pytest.mark.asyncio
async def test_publication_sse_replays_cursor_and_accepts_later_jobs(
    tmp_path,
    monkeypatch,
) -> None:
    store = PublicationEventStore("publication-sse", runs_dir=tmp_path)
    first = store.append(
        "publication.queued",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "queued", "attempt": 0},
        dedupe_key="pub-1:queued",
    )
    completed = store.append(
        "publication.completed",
        publication_id="pub-1",
        payload={"format": "pdf", "status": "completed", "attempt": 1},
        dedupe_key="pub-1:completed",
    )
    real_from_config = server.Configuration.from_runnable_config

    def fast_sse_config(config):
        resolved = real_from_config(config)
        resolved.sse_poll_interval_ms = 1
        return resolved

    monkeypatch.setattr(
        server.Configuration,
        "from_runnable_config",
        fast_sse_config,
    )
    monkeypatch.setattr(
        server,
        "get_publisher_settings",
        lambda: PublisherSettings(runs_dir=tmp_path, sse_idle_seconds=1),
    )

    iterator = server._publication_event_iterator(  # noqa: SLF001
        store,
        after=first.sequence,
    )
    assert "event: publication.completed" in await anext(iterator)

    later = store.append(
        "publication.queued",
        publication_id="pub-2",
        payload={"format": "docx", "status": "queued", "attempt": 0},
        dedupe_key="pub-2:queued",
    )
    replayed = await asyncio.wait_for(anext(iterator), timeout=0.25)
    await iterator.aclose()

    assert f"id: {later.sequence}" in replayed
    assert "event: publication.queued" in replayed
    assert completed.sequence == 2


@pytest.mark.asyncio
async def test_publication_sse_heartbeats_then_closes_after_idle_timeout(
    tmp_path,
    monkeypatch,
) -> None:
    store = PublicationEventStore("publication-heartbeat", runs_dir=tmp_path)
    real_from_config = server.Configuration.from_runnable_config

    def fast_sse_config(config):
        resolved = real_from_config(config)
        resolved.sse_poll_interval_ms = 1
        resolved.sse_heartbeat_seconds = 0.005
        return resolved

    monkeypatch.setattr(
        server.Configuration,
        "from_runnable_config",
        fast_sse_config,
    )
    monkeypatch.setattr(
        server,
        "get_publisher_settings",
        lambda: PublisherSettings(runs_dir=tmp_path, sse_idle_seconds=0.02),
    )

    iterator = server._publication_event_iterator(store)  # noqa: SLF001
    heartbeat = await asyncio.wait_for(anext(iterator), timeout=0.25)

    assert heartbeat == ": keep-alive\n\n"
    for _ in range(10):
        try:
            next_heartbeat = await asyncio.wait_for(anext(iterator), timeout=0.25)
        except StopAsyncIteration:
            break
        assert next_heartbeat == ": keep-alive\n\n"
    else:
        pytest.fail("publication SSE did not close after its idle timeout")


def test_owner_can_queue_poll_and_download_publication(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("RUNS_DIR", str(tmp_path))
    monkeypatch.setenv("PUBLISHER_ENABLED", "true")
    _completed_run(tmp_path)
    server._runs.clear()
    server.app.dependency_overrides[get_current_user] = lambda: research_principal(
        "user-1"
    )
    client = TestClient(server.app, raise_server_exceptions=False)
    try:
        queued_response = client.post(
            "/runs/publish-api/publications",
            json={"format": "pdf"},
        )
        assert queued_response.status_code == 202
        publication_id = queued_response.json()["publication_id"]
        assert queued_response.json()["status"] == "queued"
        assert "relative_path" not in queued_response.text
        queued_replay = client.post(
            "/runs/publish-api/publications",
            json={"format": "pdf"},
        )
        assert queued_replay.status_code == 202
        assert queued_replay.json()["publication_id"] == publication_id
        assert queued_replay.json()["reused"] is True
        not_ready = client.get(
            f"/runs/publish-api/publications/{publication_id}/download"
        )
        assert not_ready.status_code == 409

        settings = PublisherSettings(runs_dir=tmp_path, max_concurrent_jobs=1)
        asyncio.run(
            PublisherWorker(settings, worker_id="publisher-api-test").run_once()
        )

        status_response = client.get(
            f"/runs/publish-api/publications/{publication_id}"
        )
        assert status_response.status_code == 200
        assert status_response.json()["status"] == "completed"
        download = client.get(
            f"/runs/publish-api/publications/{publication_id}/download"
        )
        assert download.status_code == 200
        assert download.headers["content-type"] == "application/pdf"
        assert download.headers["x-content-type-options"] == "nosniff"
        artifact_sha256 = status_response.json()["artifact"]["sha256"]
        assert download.headers["etag"] == f'"{artifact_sha256}"'
        assert download.headers["content-disposition"].startswith("attachment;")
        assert download.content.startswith(b"%PDF")
        completed_replay = client.post(
            "/runs/publish-api/publications",
            json={"format": "pdf"},
        )
        assert completed_replay.status_code == 200
        assert completed_replay.json()["publication_id"] == publication_id
        assert completed_replay.json()["reused"] is True

        snapshot = client.get("/runs/publish-api").json()
        assert snapshot["output"]["preferred_output_format"] == "pdf"
        assert snapshot["output"]["publications"][0]["status"] == "completed"
    finally:
        server.app.dependency_overrides.clear()


def test_other_owner_cannot_list_publications(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("RUNS_DIR", str(tmp_path))
    _completed_run(tmp_path, "private-publication")
    server._runs.clear()
    server.app.dependency_overrides[get_current_user] = lambda: research_principal(
        "user-2"
    )
    client = TestClient(server.app, raise_server_exceptions=False)
    try:
        response = client.get("/runs/private-publication/publications")
    finally:
        server.app.dependency_overrides.clear()

    assert response.status_code == 404


def test_create_publication_checks_owner_before_publisher_availability(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setenv("RUNS_DIR", str(tmp_path))
    monkeypatch.setenv("PUBLISHER_ENABLED", "false")
    _completed_run(tmp_path, "private-disabled-publication")
    server._runs.clear()
    server.app.dependency_overrides[get_current_user] = lambda: research_principal(
        "user-2"
    )
    client = TestClient(server.app, raise_server_exceptions=False)
    try:
        response = client.post(
            "/runs/private-disabled-publication/publications",
            json={"format": "pdf"},
        )
    finally:
        server.app.dependency_overrides.clear()

    assert response.status_code == 404


def test_owner_can_retry_transient_failure_after_attempts_are_exhausted(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setenv("RUNS_DIR", str(tmp_path))
    monkeypatch.setenv("PUBLISHER_ENABLED", "true")
    monkeypatch.setenv("PUBLISHER_MAX_ATTEMPTS", "1")
    _completed_run(tmp_path, "retry-publication")
    server._runs.clear()
    server.app.dependency_overrides[get_current_user] = lambda: research_principal(
        "user-1"
    )
    client = TestClient(server.app, raise_server_exceptions=False)
    try:
        queued = client.post(
            "/runs/retry-publication/publications",
            json={"format": "pdf"},
        )
        assert queued.status_code == 202
        publication_id = queued.json()["publication_id"]
        store = PublicationJobStore("retry-publication", runs_dir=tmp_path)
        claimed = store.claim(
            publication_id,
            worker_id="failed-worker",
            lease_seconds=30,
        )
        assert claimed is not None
        failed, requeued = store.fail(
            publication_id,
            worker_id="failed-worker",
            error_code="publication_io_failed",
            retryable=True,
        )
        assert requeued is False
        assert failed.status == "failed"
        assert failed.retryable is True
        failed_replay = client.post(
            "/runs/retry-publication/publications",
            json={"format": "pdf"},
        )
        assert failed_replay.status_code == 200
        assert failed_replay.json()["status"] == "failed"
        assert failed_replay.json()["reused"] is True

        retried = client.post(
            f"/runs/retry-publication/publications/{publication_id}/retry"
        )
        assert retried.status_code == 202
        assert retried.json()["status"] == "queued"
        assert retried.json()["max_attempts"] == 2
        assert PublicationEventStore(
            "retry-publication",
            runs_dir=tmp_path,
        ).read()[-1].type == "publication.requeued"

        processed = asyncio.run(
            PublisherWorker(
                PublisherSettings(runs_dir=tmp_path, max_concurrent_jobs=1),
                worker_id="retry-worker",
            ).run_once()
        )
        assert processed == 1
        assert store.get(publication_id).status == "completed"
    finally:
        server.app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_retry_toctou_missing_job_returns_404(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("RUNS_DIR", str(tmp_path))
    context = _completed_run(tmp_path, "retry-toctou")
    store = PublicationJobStore("retry-toctou", runs_dir=tmp_path)
    markdown = context.brief_path.parent.joinpath("final_report.md").read_text()
    job, _created = store.enqueue(
        report_sha256=hashlib.sha256(markdown.encode()).hexdigest(),
        publication_format="pdf",
        theme=PublisherTheme(locale="en-US"),
        max_attempts=1,
    )
    claimed = store.claim(job.publication_id, worker_id="failed-worker", lease_seconds=30)
    assert claimed is not None
    store.fail(
        job.publication_id,
        worker_id="failed-worker",
        error_code="publication_io_failed",
        retryable=True,
    )
    server._runs.clear()

    def disappear_during_retry(_self, _publication_id, **_kwargs):
        raise FileNotFoundError(job.publication_id)

    monkeypatch.setattr(PublicationJobStore, "retry", disappear_during_retry)

    with pytest.raises(HTTPException) as error:
        await server.retry_publication(
            "retry-toctou",
            job.publication_id,
            research_principal("user-1"),
        )
    assert error.value.status_code == 404
    assert error.value.detail == "publication_not_found"


def test_worker_emits_failure_when_expired_lease_exhausts_attempts(
    tmp_path,
    monkeypatch,
) -> None:
    from open_deep_research.report import publication_store as store_module

    context = _completed_run(tmp_path, "expired-publication")
    store = PublicationJobStore("expired-publication", runs_dir=tmp_path)
    report = context.brief_path.parent.joinpath("final_report.md").read_text()
    job, _created = store.enqueue(
        report_sha256=hashlib.sha256(report.encode()).hexdigest(),
        publication_format="pdf",
        theme=PublisherTheme(locale="en-US"),
        max_attempts=1,
    )
    now = 100.0
    monkeypatch.setattr(store_module.time, "time", lambda: now)
    claimed = store.claim(
        job.publication_id,
        worker_id="crashed-worker",
        lease_seconds=1,
    )
    assert claimed is not None
    now = 102.0

    processed = asyncio.run(
        PublisherWorker(
            PublisherSettings(
                runs_dir=tmp_path,
                lease_seconds=1,
                max_concurrent_jobs=1,
            ),
            worker_id="recovery-worker",
        ).run_once()
    )

    assert processed == 0
    failed = store.get(job.publication_id)
    assert failed is not None
    assert failed.status == "failed"
    assert failed.retryable is True
    assert failed.error_code == "publication_attempts_exhausted"
    events = PublicationEventStore(
        "expired-publication",
        runs_dir=tmp_path,
    ).read()
    assert events[-1].type == "publication.failed"
    assert events[-1].payload["retryable"] is True


def test_worker_reconciles_terminal_event_after_post_commit_crash(tmp_path, monkeypatch) -> None:
    """A restart repairs the state/event gap left after artifact commit."""
    context = _completed_run(tmp_path, "terminal-event-recovery")
    store = PublicationJobStore("terminal-event-recovery", runs_dir=tmp_path)
    markdown = context.brief_path.parent.joinpath("final_report.md").read_text()
    job, _created = store.enqueue(
        report_sha256=hashlib.sha256(markdown.encode()).hexdigest(),
        publication_format="json",
        theme=PublisherTheme(locale="en-US"),
        max_attempts=1,
    )
    worker = PublisherWorker(
        PublisherSettings(runs_dir=tmp_path, max_concurrent_jobs=1),
        worker_id="recovery-worker",
    )

    async def drop_events(*_args, **_kwargs):
        """Simulate a process exit around the normal event append calls."""

    monkeypatch.setattr(worker, "_publish_event", drop_events)
    assert asyncio.run(worker.run_once()) == 1

    assert PublicationEventStore(
        "terminal-event-recovery",
        runs_dir=tmp_path,
    ).read() == []
    restarted = PublisherWorker(
        PublisherSettings(runs_dir=tmp_path, max_concurrent_jobs=1),
        worker_id="restarted-worker",
    )
    assert asyncio.run(restarted.run_once()) == 0

    events = PublicationEventStore(
        "terminal-event-recovery",
        runs_dir=tmp_path,
    ).read()
    assert [event.type for event in events] == ["publication.completed"]
    assert events[0].payload["download_url"].endswith("/download")
    assert store.get(job.publication_id).status == "completed"


def test_worker_reconciles_terminal_events_only_once_per_instance(
    tmp_path,
    monkeypatch,
) -> None:
    from open_deep_research.report import publisher_worker as worker_module

    calls = 0

    def count_reconcile(_settings) -> int:
        nonlocal calls
        calls += 1
        return 0

    monkeypatch.setattr(worker_module, "_reconcile_terminal_events", count_reconcile)
    worker = PublisherWorker(PublisherSettings(runs_dir=tmp_path), worker_id="once-worker")

    async def run_twice() -> None:
        await worker.run_once()
        await worker.run_once()

    asyncio.run(run_twice())
    assert calls == 1


def test_commit_reuses_identical_artifact_without_replace(tmp_path, monkeypatch) -> None:
    from open_deep_research.report import publication_store as store_module

    context = _completed_run(tmp_path, "reuse-artifact")
    store = PublicationJobStore("reuse-artifact", runs_dir=tmp_path)
    markdown = context.brief_path.parent.joinpath("final_report.md").read_text()
    job, _created = store.enqueue(
        report_sha256=hashlib.sha256(markdown.encode()).hexdigest(),
        publication_format="markdown",
        theme=PublisherTheme(locale="en-US"),
        max_attempts=2,
    )
    claimed = store.claim(job.publication_id, worker_id="reuse-worker", lease_seconds=30)
    assert claimed is not None
    rendered = RenderedArtifact(
        content=b"deterministic",
        media_type="text/markdown",
        extension="md",
    )
    first = store.commit_file(
        claimed,
        rendered,
        report_title="Reuse",
        max_output_bytes=100,
        worker_id="reuse-worker",
    )

    def fail_replace(_source, _target):
        raise PermissionError("download holds file open")

    monkeypatch.setattr(store_module.os, "replace", fail_replace)
    second = store.commit_file(
        claimed,
        rendered,
        report_title="Reuse",
        max_output_bytes=100,
        worker_id="reuse-worker",
    )

    assert second.sha256 == first.sha256
    assert second.size_bytes == first.size_bytes


def test_stale_worker_cannot_overwrite_newer_lease_artifact(tmp_path, monkeypatch) -> None:
    """A fenced renderer cannot replace bytes committed by a newer worker."""
    from open_deep_research.report import publication_store as store_module

    context = _completed_run(tmp_path, "lease-fence")
    store = PublicationJobStore("lease-fence", runs_dir=tmp_path)
    markdown = context.brief_path.parent.joinpath("final_report.md").read_text()
    job, _created = store.enqueue(
        report_sha256=hashlib.sha256(markdown.encode()).hexdigest(),
        publication_format="markdown",
        theme=PublisherTheme(locale="en-US"),
        max_attempts=3,
    )
    clock = {"now": 100.0}
    monkeypatch.setattr(store_module.time, "time", lambda: clock["now"])
    first = store.claim(job.publication_id, worker_id="worker-a", lease_seconds=1)
    assert first is not None
    clock["now"] = 102.0
    second = store.claim(job.publication_id, worker_id="worker-b", lease_seconds=1)
    assert second is not None

    newer = store.commit_file(
        second,
        RenderedArtifact(
            content=b"newer",
            media_type="text/markdown",
            extension="md",
        ),
        report_title="Lease fence",
        max_output_bytes=100,
        worker_id="worker-b",
    )
    store.complete(job.publication_id, worker_id="worker-b", artifact=newer)

    with pytest.raises(RuntimeError, match="publication_lease_lost"):
        store.commit_file(
            first,
            RenderedArtifact(
                content=b"stale",
                media_type="text/markdown",
                extension="md",
            ),
            report_title="Lease fence",
            max_output_bytes=100,
            worker_id="worker-a",
        )
    completed = store.get(job.publication_id)
    assert completed is not None
    assert store.artifact_path(completed).read_bytes() == b"newer"
