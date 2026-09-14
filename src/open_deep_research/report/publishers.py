"""Deterministic file publishers for canonical reports."""

from __future__ import annotations

import hashlib
import html
import io
import json
import re
import textwrap
import unicodedata
from collections.abc import Iterable
from enum import Enum
from typing import Any, Protocol

import fitz  # type: ignore[import-untyped]
from docx import Document  # type: ignore[import-untyped]
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt
from docx.shared import RGBColor as DocxRGBColor
from pptx import Presentation  # type: ignore[import-untyped]
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt

from .canonical import runs_to_plain_text, validate_canonical_report
from .models import (
    CanonicalReport,
    CodeBlock,
    InlineRun,
    ListBlock,
    ParagraphBlock,
    PublisherTheme,
    QuoteBlock,
    RenderedArtifact,
    ReportBlock,
    TableBlock,
)

PDF_MEDIA_TYPE = "application/pdf"
DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PPTX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
JSON_MEDIA_TYPE = "application/json"
MARKDOWN_MEDIA_TYPE = "text/markdown"

FORMAT_ALIASES = {
    "structured_json": "json",
    "slides": "pptx",
}
PUBLICATION_FORMATS = frozenset(
    {"markdown", "json", "pdf", "docx", "pptx", "one_pager"}
)
_MARKDOWN_ESCAPE_RE = re.compile(r"([\\`*{}\[\]()<>#+\-.!_|>])")
_BACKTICK_RUN_RE = re.compile(r"`+")
_CODE_LANGUAGE_RE = re.compile(r"[^A-Za-z0-9_+-]")


class PublicationRenderError(RuntimeError):
    """A stable publisher failure suitable for a public job error code."""

    def __init__(self, code: str, *, retryable: bool = False) -> None:
        """Initialize a normalized renderer failure."""
        super().__init__(code)
        self.code = code
        self.retryable = retryable


class Publisher(Protocol):
    """Render one canonical report into a validated byte artifact."""

    format: str

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Return one complete artifact."""
        ...


def resolve_publication_format(value: str) -> str:
    """Return a canonical publication format or raise a stable error."""
    raw = value.value if isinstance(value, Enum) else str(value or "")
    normalized_value = raw.strip().lower()
    normalized = FORMAT_ALIASES.get(normalized_value, normalized_value)
    if normalized not in PUBLICATION_FORMATS:
        raise ValueError("unsupported_publication_format")
    return normalized


def _font_names(theme: PublisherTheme) -> tuple[str, str, str]:
    if theme.font_family == "serif":
        return (
            "'Times New Roman', 'Noto Serif CJK SC', serif",
            "Times New Roman",
            "Times New Roman",
        )
    if theme.font_family == "sans":
        return "Arial, 'Noto Sans CJK SC', sans-serif", "Arial", "Arial"
    return "'Noto Sans CJK SC', sans-serif", "Noto Sans CJK SC", "Noto Sans CJK SC"


def _hex_rgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[index : index + 2], 16) for index in (1, 3, 5))  # type: ignore[return-value]


def _html_text(value: str, *, width: int = 60, force_wrap: bool = False) -> str:
    """Escape text while preserving line breaks and breaking long tokens."""
    normalized = str(value).replace("\r\n", "\n").replace("\r", "\n")
    rendered_lines: list[str] = []
    for line in normalized.split("\n"):
        long_token = max(
            (len(token) for token in re.split(r"\s+", line) if token),
            default=0,
        ) > width
        if force_wrap or long_token:
            chunks = [
                line[offset : offset + width]
                for offset in range(0, len(line), width)
            ] or [""]
        else:
            chunks = [line]
        rendered_lines.append("<br>".join(html.escape(chunk) for chunk in chunks))
    return "<br>".join(rendered_lines)


def _heading_html(
    value: str,
    *,
    level: int,
    compact: bool = False,
) -> str:
    """Escape a heading with a width suited to its rendered font size."""
    if compact:
        width = {1: 22, 2: 36}.get(level, 44)
    else:
        width = {1: 18, 2: 28}.get(level, 36)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    wrapped: list[str] = []
    for line in text.split("\n"):
        wrapped.extend(
            textwrap.wrap(
                line,
                width=width,
                break_long_words=True,
                break_on_hyphens=False,
                replace_whitespace=False,
                drop_whitespace=True,
            )
            or [""]
        )
    return "<br>".join(html.escape(line) for line in wrapped)


def _runs_html(runs: Iterable[InlineRun]) -> str:
    rendered: list[str] = []
    for run in runs:
        value = _html_text(run.text)
        if run.code:
            value = f"<code>{value}</code>"
        if run.italic:
            value = f"<em>{value}</em>"
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.href:
            value = f'<a href="{html.escape(run.href, quote=True)}">{value}</a>'
        rendered.append(value)
    return "".join(rendered)


def _safe_code_language(value: str | None) -> str:
    """Return a bounded Markdown/HTML code-language token."""
    return _CODE_LANGUAGE_RE.sub("", value or "")[:80]


def _block_html(block: ReportBlock) -> str:
    if isinstance(block, ParagraphBlock):
        return f"<p>{_runs_html(block.runs)}</p>"
    if isinstance(block, ListBlock):
        tag = "ol" if block.ordered else "ul"
        start = f' start="{block.start}"' if block.ordered and block.start != 1 else ""
        items = "".join(f"<li>{_runs_html(item)}</li>" for item in block.items)
        return f"<{tag}{start}>{items}</{tag}>"
    if isinstance(block, TableBlock):
        head = ""
        if block.headers:
            head = "<thead><tr>" + "".join(
                f"<th>{_runs_html(cell)}</th>" for cell in block.headers
            ) + "</tr></thead>"
        rows = "".join(
            "<tr>"
            + "".join(f"<td>{_runs_html(cell)}</td>" for cell in row)
            + "</tr>"
            for row in block.rows
        )
        return f"<table>{head}<tbody>{rows}</tbody></table>"
    if isinstance(block, QuoteBlock):
        return f"<blockquote>{_runs_html(block.runs)}</blockquote>"
    if isinstance(block, CodeBlock):
        language = _safe_code_language(block.language)
        wrapped_lines: list[str] = []
        normalized_code = block.text.replace("\r\n", "\n").replace("\r", "\n")
        for line in normalized_code.split("\n"):
            wrapped_lines.extend(
                line[offset : offset + 60]
                for offset in range(0, len(line), 60)
            )
            if not line:
                wrapped_lines.append("")
        # MuPDF's HTML layout can clip a single very tall ``pre`` element at a
        # page boundary.  Keep each element bounded so the Story can paginate
        # every code line instead of silently dropping the tail.
        chunks = [
            wrapped_lines[offset : offset + 30]
            for offset in range(0, len(wrapped_lines), 30)
        ] or [[]]
        return "".join(
            f'<pre data-language="{language}"><code>{html.escape(chr(10).join(chunk))}</code></pre>'
            for chunk in chunks
        )
    return ""


def _report_html(report: CanonicalReport, theme: PublisherTheme) -> str:
    labels = {
        "sources": "Sources" if theme.locale == "en-US" else "来源",
        "partial": "Partial report" if theme.locale == "en-US" else "部分报告",
    }
    body = [f"<h1>{_heading_html(report.title, level=1)}</h1>"]
    if report.completion_status == "partial":
        body.append(f'<p class="status">{labels["partial"]}</p>')
    body.extend(_block_html(block) for block in report.summary_blocks)
    for section in report.sections:
        level = min(6, max(2, section.level))
        body.append(
            f"<h{level}>{_heading_html(section.title, level=level)}</h{level}>"
        )
        body.extend(_block_html(block) for block in section.blocks)
    if report.sources:
        body.append(f'<h2>{labels["sources"]}</h2><ol class="sources">')
        body.extend(
            f'<li><a href="{html.escape(source.url, quote=True)}">'
            f"{_html_text(source.title or source.url)}</a></li>"
            for source in report.sources
        )
        body.append("</ol>")
    return "<html><body>" + "".join(body) + "</body></html>"


def _runs_markdown(runs: Iterable[InlineRun]) -> str:
    """Serialize inline runs for the Markdown publisher fallback."""
    rendered: list[str] = []
    for run in runs:
        if run.code:
            longest = max(
                (len(match.group(0)) for match in _BACKTICK_RUN_RE.finditer(run.text)),
                default=0,
            )
            fence = "`" * max(1, longest + 1)
            padding = run.text.startswith(("`", " ")) or run.text.endswith(("`", " "))
            value = f"{fence}{' ' if padding else ''}{run.text}{' ' if padding else ''}{fence}"
        else:
            value = _MARKDOWN_ESCAPE_RE.sub(r"\\\1", run.text)
        if run.bold:
            value = f"**{value}**"
        if run.italic:
            value = f"*{value}*"
        if run.href:
            destination = (
                run.href.replace("\\", "%5C").replace("<", "%3C").replace(">", "%3E")
            )
            value = f"[{value}](<{destination}>)"
        rendered.append(value)
    return "".join(rendered)


def _block_markdown(block: ReportBlock) -> str:
    if isinstance(block, ParagraphBlock):
        return _runs_markdown(block.runs)
    if isinstance(block, QuoteBlock):
        return "\n".join(f"> {line}" for line in _runs_markdown(block.runs).splitlines())
    if isinstance(block, CodeBlock):
        language = _safe_code_language(block.language)
        longest = max(
            (len(match.group(0)) for match in _BACKTICK_RUN_RE.finditer(block.text)),
            default=0,
        )
        fence = "`" * max(3, longest + 1)
        body = block.text if block.text.endswith("\n") else block.text + "\n"
        return f"{fence}{language}\n{body}{fence}"
    if isinstance(block, ListBlock):
        lines: list[str] = []
        for index, item in enumerate(block.items):
            marker = (
                f"{block.start + index}." if block.ordered else "-"
            )
            lines.append(f"{marker} {_runs_markdown(item)}")
        return "\n".join(lines)
    if isinstance(block, TableBlock):
        rows = [block.headers, *block.rows] if block.headers else list(block.rows)
        if not rows:
            return ""
        rendered_rows = [
            "| " + " | ".join(_runs_markdown(cell) for cell in row) + " |"
            for row in rows
        ]
        if block.headers:
            separator = "| " + " | ".join("---" for _ in block.headers) + " |"
            rendered_rows.insert(1, separator)
        return "\n".join(rendered_rows)
    return ""


def canonical_report_to_markdown(report: CanonicalReport) -> str:
    """Render a canonical report when the original Markdown is unavailable."""
    parts = [f"# {_MARKDOWN_ESCAPE_RE.sub(r'\\\1', report.title)}"]
    parts.extend(_block_markdown(block) for block in report.summary_blocks)
    for section in report.sections:
        title = _MARKDOWN_ESCAPE_RE.sub(r"\\\1", section.title)
        parts.append(f"{'#' * min(6, max(2, section.level))} {title}")
        parts.extend(_block_markdown(block) for block in section.blocks)
    if report.sources:
        parts.append("## Sources")
        parts.extend(
            "- ["
            + _MARKDOWN_ESCAPE_RE.sub(r"\\\1", source.title or source.url)
            + "](<"
            + source.url.replace("\\", "%5C").replace("<", "%3C").replace(">", "%3E")
            + ">)"
            for source in report.sources
        )
    return "\n\n".join(part for part in parts if part).rstrip() + "\n"


def _pdf_css(theme: PublisherTheme, *, compact: bool = False) -> str:
    css_font, _, _ = _font_names(theme)
    base = "9.2pt" if compact else "10.5pt"
    return f"""
        body {{ font-family: {css_font}; font-size: {base}; color: #20252A; line-height: 1.42; }}
        h1 {{ color: {theme.primary_color}; font-size: {'22pt' if compact else '28pt'}; margin: 0 0 14pt; }}
        h2 {{ color: {theme.primary_color}; font-size: {'14pt' if compact else '18pt'}; margin: 16pt 0 7pt; }}
        h3, h4, h5, h6 {{ color: #303840; margin: 12pt 0 5pt; }}
        p {{ margin: 0 0 7pt; }}
        ul, ol {{ margin: 3pt 0 8pt 17pt; }}
        li {{ margin-bottom: 3pt; }}
        a {{ color: {theme.primary_color}; text-decoration: none; }}
        code {{ font-family: monospace; background: #F1F3F4; }}
        pre {{ font-family: monospace; font-size: 8.5pt; background: #F1F3F4; padding: 7pt; white-space: pre-wrap; overflow-wrap: anywhere; word-break: break-word; }}
        blockquote {{ border-left: 3pt solid {theme.accent_color}; padding-left: 9pt; color: #4A545C; }}
        table {{ border-collapse: collapse; width: 100%; margin: 6pt 0 10pt; font-size: {'7.8pt' if compact else '9pt'}; }}
        th {{ background: #E6ECE7; color: #20252A; font-weight: bold; }}
        th, td {{ border: 0.6pt solid #B8C0C5; padding: 4pt; vertical-align: top; }}
        .status {{ color: #8A4B08; font-weight: bold; }}
        .sources {{ font-size: 8.5pt; }}
    """


def _page_rect(theme: PublisherTheme) -> fitz.Rect:
    return fitz.paper_rect("letter" if theme.pdf_page_size == "letter" else "a4")


def _validate_pdf(content: bytes, *, max_pages: int = 100) -> int:
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise PublicationRenderError("publisher_pdf_invalid") from exc
    try:
        count = document.page_count
        if count < 1:
            raise PublicationRenderError("publisher_pdf_empty")
        if count > max_pages:
            raise PublicationRenderError("publisher_pdf_page_limit_exceeded")
        return count
    finally:
        document.close()


class MarkdownPublisher:
    """Provide a deterministic Markdown fallback for direct Registry callers."""

    format = "markdown"

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Render canonical content as UTF-8 Markdown."""
        del theme
        return RenderedArtifact(
            content=canonical_report_to_markdown(report).encode("utf-8"),
            media_type=MARKDOWN_MEDIA_TYPE,
            extension="md",
        )


class PdfPublisher:
    """Publish an accessible, paginated PDF with retained links."""

    format = "pdf"

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Render and validate a paginated PDF."""
        page_rect = _page_rect(theme)
        content_rect = fitz.Rect(
            page_rect.x0 + 48,
            page_rect.y0 + 46,
            page_rect.x1 - 48,
            page_rect.y1 - 48,
        )

        def rectfn(_rect_num: int, _filled: fitz.Rect):
            return page_rect, content_rect, fitz.Identity

        try:
            story = fitz.Story(
                _report_html(report, theme),
                user_css=_pdf_css(theme),
                em=10.5,
            )
            document = story.write_with_links(rectfn)
            footer = theme.footer_text.strip()
            for index, page in enumerate(document):
                label = f"{footer}  " if footer else ""
                label += f"{index + 1} / {document.page_count}"
                try:
                    page.insert_text(
                        fitz.Point(page_rect.x0 + 48, page_rect.y1 - 22),
                        label,
                        fontsize=8,
                        fontname="china-s" if theme.font_family == "cjk_sans" else "helv",
                        color=(0.35, 0.39, 0.42),
                    )
                except Exception:
                    page.insert_text(
                        fitz.Point(page_rect.x0 + 48, page_rect.y1 - 22),
                        str(index + 1),
                        fontsize=8,
                    )
            content = document.tobytes(garbage=4, deflate=True)
            document.close()
        except PublicationRenderError:
            raise
        except Exception as exc:
            raise PublicationRenderError("publisher_pdf_render_failed") from exc
        page_count = _validate_pdf(content)
        return RenderedArtifact(
            content=content,
            media_type=PDF_MEDIA_TYPE,
            extension="pdf",
            page_count=page_count,
        )


def _plain_block_lines(block: ReportBlock) -> list[str]:
    if isinstance(block, ParagraphBlock | QuoteBlock):
        value = runs_to_plain_text(block.runs)
        return [value] if value else []
    if isinstance(block, ListBlock):
        return [runs_to_plain_text(item) for item in block.items if runs_to_plain_text(item)]
    if isinstance(block, TableBlock):
        rows = [block.headers, *block.rows] if block.headers else list(block.rows)
        return [
            " | ".join(runs_to_plain_text(cell) for cell in row)
            for row in rows
        ]
    if isinstance(block, CodeBlock):
        return [block.text.strip()] if block.text.strip() else []
    return []


_ONE_PAGER_SECTION_GROUPS: tuple[re.Pattern[str], ...] = (
    re.compile(
        r"recommend|conclusion|decision|verdict|建议|推荐|结论|决策|判断",
        re.IGNORECASE,
    ),
    re.compile(
        r"key\s+findings?|highlights?|takeaways?|关键|要点|发现|核心",
        re.IGNORECASE,
    ),
    re.compile(
        r"limitation|uncertaint|risk|caveat|限制|不确定|风险|局限",
        re.IGNORECASE,
    ),
)


def _select_one_pager_sections(
    report: CanonicalReport,
    *,
    max_sections: int = 3,
) -> tuple[list[Any], bool]:
    """Select high-value sections without truncating a selected conclusion.

    Section selection is intentionally deterministic.  Once a section is
    selected, all of its blocks are retained; the renderer reports overflow
    instead of silently deleting the middle of a recommendation or finding.
    """
    if max_sections <= 0 or not report.sections:
        return [], bool(report.sections)
    selected_indices: list[int] = []
    for pattern in _ONE_PAGER_SECTION_GROUPS:
        for index, section in enumerate(report.sections):
            if index in selected_indices:
                continue
            if pattern.search(section.title):
                selected_indices.append(index)
                break
        if len(selected_indices) >= max_sections:
            break
    # Fill unused slots in source order so short reports still show their
    # overview/context sections.  Sorting restores the report's original order.
    for index in range(len(report.sections)):
        if len(selected_indices) >= max_sections:
            break
        if index not in selected_indices:
            selected_indices.append(index)
    selected_indices.sort()
    return [report.sections[index] for index in selected_indices], (
        len(selected_indices) < len(report.sections)
    )


def _one_pager_html(report: CanonicalReport, theme: PublisherTheme) -> str:
    selected, omitted_sections = _select_one_pager_sections(report)
    body = [f"<h1>{_heading_html(report.title, level=1, compact=True)}</h1>"]
    # Keep the summary bounded and make the omission explicit.  Unlike the
    # selected decision sections below, summary trimming is a deliberate,
    # labelled product rule rather than an accidental layout truncation.
    summary_limit = 2
    body.extend(_block_html(block) for block in report.summary_blocks[:summary_limit])
    for section in selected:
        level = min(6, max(2, section.level))
        body.append(
            f"<h{level}>{_heading_html(section.title, level=level, compact=True)}</h{level}>"
        )
        body.extend(_block_html(block) for block in section.blocks)
    condensed = omitted_sections or len(report.summary_blocks) > summary_limit
    if condensed:
        notice = (
            "Content condensed under the one-page publication rule."
            if theme.locale == "en-US"
            else "内容已按单页发布规则压缩。"
        )
        body.append(f'<p class="status">{notice}</p>')
    if report.sources:
        label = "Sources" if theme.locale == "en-US" else "来源"
        body.append(f"<h2>{label}</h2><ol class=\"sources\">")
        body.extend(
            f'<li><a href="{html.escape(source.url, quote=True)}">'
            f"{_html_text(source.title or source.url)}</a></li>"
            for source in report.sources[:5]
        )
        body.append("</ol>")
    return "<html><body>" + "".join(body) + "</body></html>"


class OnePagerPublisher:
    """Publish a deterministic, strictly single-page PDF."""

    format = "one_pager"

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Render and validate one strictly single-page PDF."""
        page_rect = _page_rect(theme)
        document = None
        try:
            document = fitz.open()
            page = document.new_page(width=page_rect.width, height=page_rect.height)
            box = fitz.Rect(38, 34, page_rect.width - 38, page_rect.height - 48)
            spare, _scale = page.insert_htmlbox(
                box,
                _one_pager_html(report, theme),
                css=_pdf_css(theme, compact=True),
                scale_low=0.78,
            )
            if spare < 0:
                raise PublicationRenderError("one_page_overflow")
            footer = theme.footer_text.strip()
            if footer:
                try:
                    page.insert_text(
                        fitz.Point(38, page_rect.height - 22),
                        footer,
                        fontsize=7.5,
                        fontname="china-s" if theme.font_family == "cjk_sans" else "helv",
                        color=(0.35, 0.39, 0.42),
                    )
                except Exception:
                    page.insert_text(
                        fitz.Point(38, page_rect.height - 22),
                        footer,
                        fontsize=7.5,
                        color=(0.35, 0.39, 0.42),
                    )
            content = document.tobytes(garbage=4, deflate=True)
        except PublicationRenderError:
            raise
        except Exception as exc:
            raise PublicationRenderError("publisher_one_pager_render_failed") from exc
        finally:
            if document is not None:
                document.close()
        page_count = _validate_pdf(content, max_pages=1)
        if page_count != 1:
            raise PublicationRenderError("one_page_overflow")
        return RenderedArtifact(
            content=content,
            media_type=PDF_MEDIA_TYPE,
            extension="pdf",
            page_count=1,
        )


def _add_docx_hyperlink(
    paragraph: Any,
    text: str,
    href: str,
    color: str,
    *,
    font_name: str = "Noto Sans CJK SC",
    bold: bool = False,
    italic: bool = False,
) -> None:
    relationship_id = paragraph.part.relate_to(
        href,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:eastAsia"), font_name)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color.removeprefix("#"))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([fonts, color_node, underline])
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # noqa: SLF001


def _add_docx_runs(paragraph: Any, runs: Iterable[InlineRun], theme: PublisherTheme) -> None:
    _, office_font, _ = _font_names(theme)
    for source in runs:
        if source.href:
            _add_docx_hyperlink(
                paragraph,
                source.text,
                source.href,
                theme.primary_color,
                font_name="Courier New" if source.code else office_font,
                bold=source.bold,
                italic=source.italic,
            )
            continue
        run = paragraph.add_run(source.text)
        run.bold = source.bold
        run.italic = source.italic
        run.font.name = "Courier New" if source.code else office_font
        rpr = run._element.get_or_add_rPr()  # noqa: SLF001
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:ascii"), run.font.name)
        fonts.set(qn("w:hAnsi"), run.font.name)
        fonts.set(qn("w:eastAsia"), run.font.name)
        if source.code:
            run.font.size = Pt(9)


def _add_docx_block(document: Any, block: ReportBlock, theme: PublisherTheme) -> None:
    if isinstance(block, ParagraphBlock):
        paragraph = document.add_paragraph()
        _add_docx_runs(paragraph, block.runs, theme)
        return
    if isinstance(block, ListBlock):
        style = "List Number" if block.ordered else "List Bullet"
        for item in block.items:
            paragraph = document.add_paragraph(style=style)
            _add_docx_runs(paragraph, item, theme)
        return
    if isinstance(block, TableBlock):
        column_count = max(
            len(block.headers),
            max((len(row) for row in block.rows), default=0),
        )
        if column_count == 0:
            return
        row_count = len(block.rows) + (1 if block.headers else 0)
        table = document.add_table(rows=row_count, cols=column_count)
        table.style = "Table Grid"
        offset = 0
        if block.headers:
            for index, cell_runs in enumerate(block.headers):
                _add_docx_runs(table.rows[0].cells[index].paragraphs[0], cell_runs, theme)
                for run in table.rows[0].cells[index].paragraphs[0].runs:
                    run.bold = True
            offset = 1
        for row_index, row in enumerate(block.rows, offset):
            for cell_index, cell_runs in enumerate(row):
                _add_docx_runs(
                    table.rows[row_index].cells[cell_index].paragraphs[0],
                    cell_runs,
                    theme,
                )
        return
    if isinstance(block, QuoteBlock):
        paragraph = document.add_paragraph(style="Intense Quote")
        _add_docx_runs(paragraph, block.runs, theme)
        return
    if isinstance(block, CodeBlock):
        paragraph = document.add_paragraph(style="No Spacing")
        run = paragraph.add_run(block.text)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)


class DocxPublisher:
    """Publish an editable Word report."""

    format = "docx"

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Render and validate one editable Word document."""
        try:
            document = Document()
            section = document.sections[0]
            if theme.pdf_page_size == "letter":
                section.page_width = Inches(8.5)
                section.page_height = Inches(11)
            else:
                section.page_width = Mm(210)
                section.page_height = Mm(297)
            section.top_margin = Mm(20)
            section.bottom_margin = Mm(20)
            section.left_margin = Mm(20)
            section.right_margin = Mm(20)
            # Office core properties cap the title at 255 characters; keep
            # the full canonical title in the document body below.
            document.core_properties.title = report.title[:255]
            _, office_font, _ = _font_names(theme)
            styles = document.styles
            styles["Normal"].font.name = office_font
            styles["Normal"].font.size = Pt(10.5)
            for style_name in ("Normal", "No Spacing", "Intense Quote", *[f"Heading {level}" for level in range(1, 7)]):
                if style_name not in styles:
                    continue
                style_rpr = styles[style_name]._element.get_or_add_rPr()  # noqa: SLF001
                style_fonts = style_rpr.rFonts
                if style_fonts is None:
                    style_fonts = OxmlElement("w:rFonts")
                    style_rpr.insert(0, style_fonts)
                style_fonts.set(qn("w:ascii"), office_font)
                style_fonts.set(qn("w:hAnsi"), office_font)
                style_fonts.set(qn("w:eastAsia"), office_font)
            for level in range(1, 7):
                style_name = f"Heading {level}"
                if style_name in styles:
                    styles[style_name].font.name = office_font
                    styles[style_name].font.color.rgb = DocxRGBColor(
                        *_hex_rgb(theme.primary_color)
                    )
            title = document.add_heading(report.title, level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for block in report.summary_blocks:
                _add_docx_block(document, block, theme)
            for report_section in report.sections:
                document.add_heading(
                    report_section.title,
                    level=min(6, max(1, report_section.level - 1)),
                )
                for block in report_section.blocks:
                    _add_docx_block(document, block, theme)
            if report.sources:
                document.add_heading(
                    "Sources" if theme.locale == "en-US" else "来源",
                    level=1,
                )
                for source in report.sources:
                    paragraph = document.add_paragraph(style="List Number")
                    _add_docx_hyperlink(
                        paragraph,
                        source.title or source.url,
                        source.url,
                        theme.primary_color,
                        font_name=office_font,
                    )
            if theme.footer_text:
                footer = section.footer.paragraphs[0]
                footer.text = theme.footer_text
                footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            stream = io.BytesIO()
            document.save(stream)
            content = stream.getvalue()
            Document(io.BytesIO(content))
        except Exception as exc:
            raise PublicationRenderError("publisher_docx_render_failed") from exc
        return RenderedArtifact(
            content=content,
            media_type=DOCX_MEDIA_TYPE,
            extension="docx",
        )


def _set_pptx_text_style(paragraph, theme: PublisherTheme, *, size: int = 20) -> None:
    _, _, pptx_font = _font_names(theme)
    paragraph.font.name = pptx_font
    paragraph.font.size = PptxPt(size)
    paragraph.font.color.rgb = RGBColor(32, 37, 42)


def _pptx_footer(slide: Any, presentation: Any, theme: PublisherTheme) -> None:
    if not theme.footer_text:
        return
    box = slide.shapes.add_textbox(
        PptxInches(0.55),
        presentation.slide_height - PptxInches(0.38),
        presentation.slide_width - PptxInches(1.1),
        PptxInches(0.22),
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = theme.footer_text
    paragraph.alignment = PP_ALIGN.RIGHT
    _set_pptx_text_style(paragraph, theme, size=8)


def _paginate_pptx_bullets(bullets: list[str]) -> list[list[str]]:
    """Wrap and paginate bullets without relying on PowerPoint auto-fit."""
    line_width = 72
    max_lines = 9
    max_bullets = 6
    pages: list[list[str]] = []
    current: list[str] = []
    current_lines = 0

    def flush() -> None:
        nonlocal current, current_lines
        if current:
            pages.append(current)
            current = []
            current_lines = 0

    for bullet in bullets or [""]:
        wrapped: list[str] = []
        for source_line in bullet.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            candidates = textwrap.wrap(
                source_line,
                width=line_width,
                break_long_words=True,
                break_on_hyphens=False,
                replace_whitespace=False,
                drop_whitespace=True,
            ) or [""]
            for candidate in candidates:
                line_chunk = ""
                line_chunk_width = 0
                for character in candidate:
                    character_width = (
                        2
                        if unicodedata.east_asian_width(character) in {"F", "W"}
                        else 1
                    )
                    if (
                        line_chunk
                        and line_chunk_width + character_width > line_width
                    ):
                        wrapped.append(line_chunk.rstrip())
                        line_chunk = ""
                        line_chunk_width = 0
                    line_chunk += character
                    line_chunk_width += character_width
                if line_chunk or not candidate:
                    wrapped.append(line_chunk.rstrip())
        offset = 0
        while offset < len(wrapped):
            if current and (
                len(current) >= max_bullets or current_lines >= max_lines
            ):
                flush()
            available = max_lines - current_lines
            chunk = wrapped[offset : offset + available]
            current.append("\n".join(chunk))
            current_lines += len(chunk)
            offset += len(chunk)
            if offset < len(wrapped):
                flush()
    flush()
    return pages or [[""]]


def _add_pptx_bullet_slide(
    presentation: Any,
    title: str,
    bullets: list[str],
    theme: PublisherTheme,
) -> dict:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    title_frame = slide.shapes.title.text_frame
    title_frame.word_wrap = True
    title_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    title_paragraph = title_frame.paragraphs[0]
    _set_pptx_text_style(title_paragraph, theme, size=26)
    title_paragraph.font.color.rgb = RGBColor(*_hex_rgb(theme.primary_color))
    frame = slide.placeholders[1].text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, bullet in enumerate(bullets or [""]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        _set_pptx_text_style(paragraph, theme, size=18)
    _pptx_footer(slide, presentation, theme)
    return {"title": title, "bullets": bullets}


def _add_pptx_bullet_slides(
    presentation: Any,
    title: str,
    bullets: list[str],
    theme: PublisherTheme,
) -> list[dict]:
    """Add one or more line-budgeted slides for a bullet collection."""
    pages = _paginate_pptx_bullets(bullets)
    total = len(pages)
    return [
        _add_pptx_bullet_slide(
            presentation,
            title if total == 1 else f"{title} ({index}/{total})",
            page,
            theme,
        )
        for index, page in enumerate(pages, start=1)
    ]


def _add_pptx_table_slides(
    presentation: Any,
    title: str,
    block: TableBlock,
    theme: PublisherTheme,
) -> list[dict]:
    rows = [list(row) for row in block.rows]
    previews: list[dict] = []
    column_count = max(
        len(block.headers),
        max((len(row) for row in rows), default=0),
    )
    if column_count == 0:
        return previews

    # Wide tables are split in both dimensions.  Keeping the chunks
    # deterministic makes retries produce the same deck and preview metadata.
    max_columns = 5
    max_rows = 6
    column_chunks = [
        list(range(start, min(start + max_columns, column_count)))
        for start in range(0, column_count, max_columns)
    ]
    row_chunks = [
        rows[start : start + max_rows]
        for start in range(0, len(rows), max_rows)
    ] or [[]]
    total_chunks = len(column_chunks) * len(row_chunks)
    chunk_number = 0
    for column_offset, columns in enumerate(column_chunks):
        for row_offset, chunk in enumerate(row_chunks):
            chunk_number += 1
            suffix = ""
            if total_chunks > 1:
                suffix = f" ({chunk_number}/{total_chunks})"
            slide_title = title + suffix
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = slide_title
            title_frame = slide.shapes.title.text_frame
            title_frame.word_wrap = True
            title_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            title_paragraph = title_frame.paragraphs[0]
            _set_pptx_text_style(title_paragraph, theme, size=24)
            title_paragraph.font.color.rgb = RGBColor(*_hex_rgb(theme.primary_color))
            table_rows = len(chunk) + (1 if block.headers else 0)
            table = slide.shapes.add_table(
                table_rows,
                len(columns),
                PptxInches(0.55),
                PptxInches(1.35),
                presentation.slide_width - PptxInches(1.1),
                presentation.slide_height - PptxInches(1.9),
            ).table
            row_index = 0
            if block.headers:
                for cell_index, source_index in enumerate(columns):
                    cell_runs = (
                        block.headers[source_index]
                        if source_index < len(block.headers)
                        else []
                    )
                    table.cell(0, cell_index).text = runs_to_plain_text(cell_runs)
                    fill = table.cell(0, cell_index).fill
                    fill.solid()
                    fill.fore_color.rgb = RGBColor(*_hex_rgb(theme.primary_color))
                    for paragraph in table.cell(0, cell_index).text_frame.paragraphs:
                        _set_pptx_text_style(paragraph, theme, size=11)
                        paragraph.font.color.rgb = RGBColor(255, 255, 255)
                        paragraph.font.bold = True
                row_index = 1
            for source_row in chunk:
                for cell_index, source_index in enumerate(columns):
                    cell_runs = source_row[source_index] if source_index < len(source_row) else []
                    table.cell(row_index, cell_index).text = runs_to_plain_text(cell_runs)
                    for paragraph in table.cell(row_index, cell_index).text_frame.paragraphs:
                        _set_pptx_text_style(paragraph, theme, size=10)
                row_index += 1
            _pptx_footer(slide, presentation, theme)
            preview_rows = []
            if block.headers:
                preview_rows.append(
                    [
                        runs_to_plain_text(block.headers[index])
                        if index < len(block.headers)
                        else ""
                        for index in columns
                    ]
                )
            preview_rows.extend(
                [
                    runs_to_plain_text(source_row[index])
                    if index < len(source_row)
                    else ""
                    for index in columns
                ]
                for source_row in chunk
            )
            previews.append({"title": slide_title, "table": preview_rows})
    return previews


class PptxPublisher:
    """Publish a 16:9 or 4:3 presentation with deterministic slide splitting."""

    format = "pptx"

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Render and validate one PowerPoint deck."""
        try:
            presentation = Presentation()
            slide_height = PptxInches(7.5)
            presentation.slide_height = slide_height
            if theme.pptx_aspect_ratio == "4:3":
                presentation.slide_width = slide_height * 4 // 3
            else:
                presentation.slide_width = slide_height * 16 // 9
            preview: list[dict] = []
            title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            title_slide.shapes.title.text = report.title
            title_frame = title_slide.shapes.title.text_frame
            title_frame.word_wrap = True
            title_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            title_paragraph = title_frame.paragraphs[0]
            _set_pptx_text_style(title_paragraph, theme, size=30)
            title_paragraph.font.color.rgb = RGBColor(*_hex_rgb(theme.primary_color))
            subtitle = title_slide.placeholders[1]
            subtitle.text = report.generated_at.date().isoformat()
            _set_pptx_text_style(subtitle.text_frame.paragraphs[0], theme, size=14)
            _pptx_footer(title_slide, presentation, theme)
            preview.append({"title": report.title, "bullets": []})

            summary_label = "Executive Summary" if theme.locale == "en-US" else "摘要"
            summary_pending: list[str] = []

            def flush_summary_bullets() -> None:
                """Write pending summary blocks in the standard six-item pages."""
                if not summary_pending:
                    return
                preview.extend(
                    _add_pptx_bullet_slides(
                        presentation,
                        summary_label,
                        summary_pending,
                        theme,
                    )
                )
                summary_pending.clear()

            for block in report.summary_blocks:
                if isinstance(block, TableBlock):
                    flush_summary_bullets()
                    preview.extend(
                        _add_pptx_table_slides(
                            presentation,
                            summary_label,
                            block,
                            theme,
                        )
                    )
                else:
                    summary_pending.extend(_plain_block_lines(block))
            flush_summary_bullets()
            for section in report.sections:
                pending: list[str] = []
                for block in section.blocks:
                    if isinstance(block, TableBlock):
                        if pending:
                            preview.extend(
                                _add_pptx_bullet_slides(
                                    presentation,
                                    section.title,
                                    pending,
                                    theme,
                                )
                            )
                            pending.clear()
                        preview.extend(
                            _add_pptx_table_slides(
                                presentation,
                                section.title,
                                block,
                                theme,
                            )
                        )
                    else:
                        pending.extend(_plain_block_lines(block))
                if pending or not section.blocks:
                    preview.extend(
                        _add_pptx_bullet_slides(
                            presentation,
                            section.title,
                            pending,
                            theme,
                        )
                    )
            if report.sources:
                label = "Sources" if theme.locale == "en-US" else "来源"
                source_lines = [
                    f"{source.title or 'Source'}: {source.url}"
                    for source in report.sources
                ]
                preview.extend(
                    _add_pptx_bullet_slides(
                        presentation,
                        label,
                        source_lines,
                        theme,
                    )
                )
            stream = io.BytesIO()
            presentation.save(stream)
            content = stream.getvalue()
            validated = Presentation(io.BytesIO(content))
            slide_count = len(validated.slides)
        except PublicationRenderError:
            raise
        except Exception as exc:
            raise PublicationRenderError("publisher_pptx_render_failed") from exc
        return RenderedArtifact(
            content=content,
            media_type=PPTX_MEDIA_TYPE,
            extension="pptx",
            slide_count=slide_count,
            preview={"slides": preview},
        )


class JsonPublisher:
    """Publish the full canonical model as UTF-8 JSON."""

    format = "json"

    def render(
        self,
        report: CanonicalReport,
        theme: PublisherTheme,
    ) -> RenderedArtifact:
        """Serialize the complete canonical report as stable JSON."""
        del theme
        content = json.dumps(
            report.model_dump(mode="json"),
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        ).encode("utf-8")
        return RenderedArtifact(
            content=content,
            media_type=JSON_MEDIA_TYPE,
            extension="json",
        )


PUBLISHERS: dict[str, Publisher] = {
    "markdown": MarkdownPublisher(),
    "pdf": PdfPublisher(),
    "docx": DocxPublisher(),
    "pptx": PptxPublisher(),
    "json": JsonPublisher(),
    "one_pager": OnePagerPublisher(),
}


class PublisherRegistry:
    """Registry facade used by workers and external integrations."""

    def __init__(self, publishers: dict[str, Publisher] | None = None) -> None:
        """Initialize a registry from the built-in publisher map."""
        source = publishers or PUBLISHERS
        self._publishers = {
            resolve_publication_format(key): publisher
            for key, publisher in source.items()
        }

    def get(self, publication_format: str) -> Publisher:
        """Return the publisher for a canonical or aliased format."""
        resolved = resolve_publication_format(publication_format)
        try:
            return self._publishers[resolved]
        except KeyError as exc:
            raise ValueError("unsupported_publication_format") from exc

    def formats(self) -> tuple[str, ...]:
        """Return all registered publication formats."""
        return tuple(sorted(self._publishers))


DEFAULT_PUBLISHER_REGISTRY = PublisherRegistry()


def get_publisher(publication_format: str) -> Publisher:
    """Return a built-in publisher by canonical or legacy format name."""
    return DEFAULT_PUBLISHER_REGISTRY.get(publication_format)


def render_publication(
    report: CanonicalReport,
    publication_format: str,
    theme: PublisherTheme,
    *,
    source_markdown: str | None = None,
) -> RenderedArtifact:
    """Render one canonical format, including the Markdown passthrough."""
    try:
        validate_canonical_report(report)
    except ValueError as exc:
        raise PublicationRenderError(str(exc)) from exc
    resolved = resolve_publication_format(publication_format)
    if resolved == "markdown":
        if source_markdown is None:
            return get_publisher("markdown").render(report, theme)
        if (
            hashlib.sha256(source_markdown.encode("utf-8")).hexdigest()
            != report.source_markdown_sha256
        ):
            raise PublicationRenderError("publisher_markdown_mismatch")
        return RenderedArtifact(
            content=source_markdown.encode("utf-8"),
            media_type=MARKDOWN_MEDIA_TYPE,
            extension="md",
        )
    publisher = get_publisher(resolved)
    return publisher.render(report, theme)


def validate_rendered_artifact(
    report: CanonicalReport,
    publication_format: str,
    rendered: RenderedArtifact,
    *,
    source_markdown: str | None = None,
    max_pdf_pages: int = 100,
    max_pptx_slides: int = 64,
) -> None:
    """Validate renderer bytes and metadata before they enter the queue.

    Built-in publishers validate themselves as well, but the worker performs
    this second boundary check so a swapped/custom publisher cannot commit a
    mismatched extension or an invalid package.
    """
    resolved = resolve_publication_format(publication_format)
    expected_extension = {
        "markdown": "md",
        "json": "json",
        "pdf": "pdf",
        "docx": "docx",
        "pptx": "pptx",
        "one_pager": "pdf",
    }[resolved]
    expected_media_type = {
        "markdown": MARKDOWN_MEDIA_TYPE,
        "json": JSON_MEDIA_TYPE,
        "pdf": PDF_MEDIA_TYPE,
        "docx": DOCX_MEDIA_TYPE,
        "pptx": PPTX_MEDIA_TYPE,
        "one_pager": PDF_MEDIA_TYPE,
    }[resolved]
    if rendered.extension.lower() != expected_extension or rendered.media_type != expected_media_type:
        raise PublicationRenderError("publisher_output_metadata_invalid")
    if not rendered.content:
        raise PublicationRenderError("publication_output_empty")
    try:
        if resolved == "markdown":
            if source_markdown is None:
                rendered.content.decode("utf-8")
            elif rendered.content != source_markdown.encode("utf-8"):
                raise PublicationRenderError("publisher_markdown_mismatch")
        elif resolved == "json":
            candidate = CanonicalReport.model_validate_json(rendered.content)
            if (
                candidate.run_id != report.run_id
                or candidate.source_markdown_sha256
                != report.source_markdown_sha256
                or candidate.model_dump(mode="json")
                != report.model_dump(mode="json")
            ):
                raise PublicationRenderError("publisher_json_mismatch")
        elif resolved in {"pdf", "one_pager"}:
            page_count = _validate_pdf(rendered.content, max_pages=max_pdf_pages)
            if (
                rendered.page_count is not None
                and rendered.page_count != page_count
            ):
                raise PublicationRenderError("publisher_page_count_mismatch")
            if resolved == "one_pager" and page_count != 1:
                raise PublicationRenderError("one_page_overflow")
        elif resolved == "docx":
            Document(io.BytesIO(rendered.content))
        elif resolved == "pptx":
            deck = Presentation(io.BytesIO(rendered.content))
            if len(deck.slides) < 1:
                raise PublicationRenderError("publisher_pptx_empty")
            if len(deck.slides) > max_pptx_slides:
                raise PublicationRenderError("pptx_slide_limit_exceeded")
            if (
                rendered.slide_count is not None
                and rendered.slide_count != len(deck.slides)
            ):
                raise PublicationRenderError("publisher_slide_count_mismatch")
    except PublicationRenderError:
        raise
    except Exception as exc:
        raise PublicationRenderError("publisher_output_invalid") from exc
