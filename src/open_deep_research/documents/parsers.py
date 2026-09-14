"""Deterministic document parsers with local or remote PaddleOCR support."""

from __future__ import annotations

import base64
import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import fitz
import httpx
from charset_normalizer import from_bytes

from .settings import DocumentSettings


class DocumentParseError(RuntimeError):
    """Raised when a supported file cannot produce safe, usable text."""


@dataclass(frozen=True, slots=True)
class ParsedUnit:
    """One source-located structural unit before chunking."""

    locator: str
    text: str
    heading: str | None = None


@dataclass(slots=True)
class ParsedDocument:
    """Normalized parser result persisted by the ingestion worker."""

    units: list[ParsedUnit] = field(default_factory=list)
    page_count: int | None = None
    ocr_pages: int = 0
    quality_flags: list[str] = field(default_factory=list)


_ocr_engine: Any | None = None


def _paddle_engine() -> Any:
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from paddleocr import PaddleOCR
        except ImportError as exc:
            raise DocumentParseError("document_ocr_runtime_unavailable") from exc
        try:
            _ocr_engine = PaddleOCR(
                lang="ch",
                use_doc_orientation_classify=True,
                use_doc_unwarping=False,
                use_textline_orientation=True,
            )
        except TypeError:
            _ocr_engine = PaddleOCR(lang="ch", use_angle_cls=True, show_log=False)
    return _ocr_engine


def _texts_from_ocr_result(result: Any) -> list[tuple[str, float]]:
    """Extract recognized lines from local PaddleOCR or HTTP response shapes."""
    texts: list[tuple[str, float]] = []
    values = result if isinstance(result, list) else [result]
    for item in values:
        payload = getattr(item, "json", item)
        if callable(payload):
            payload = payload()
        if isinstance(payload, str):
            try:
                payload = json.loads(payload)
            except json.JSONDecodeError:
                continue
        if isinstance(payload, dict):
            direct_text = (
                payload.get("text")
                or payload.get("transcription")
                or payload.get("rec_text")
            )
            if direct_text:
                try:
                    direct_score = float(
                        payload.get("score", payload.get("confidence", 1.0))
                    )
                except (TypeError, ValueError):
                    direct_score = 0.0
                texts.append((str(direct_text), direct_score))
            for nested_key in (
                "result",
                "results",
                "data",
                "ocr_result",
                "ocrResults",
                "ocr_results",
                "prunedResult",
                "pruned_result",
                "res",
            ):
                nested = payload.get(nested_key)
                if nested is not None:
                    texts.extend(_texts_from_ocr_result(nested))
            rec_texts = payload.get("rec_texts") or payload.get("texts") or []
            rec_scores = payload.get("rec_scores") or payload.get("scores") or []
            if isinstance(rec_texts, str):
                rec_texts = [rec_texts]
            if not isinstance(rec_scores, list | tuple):
                rec_scores = []
            for index, text in enumerate(rec_texts):
                try:
                    score = float(
                        rec_scores[index] if index < len(rec_scores) else 1.0
                    )
                except (TypeError, ValueError):
                    score = 0.0
                texts.append((str(text), score))
            continue
        if not isinstance(payload, list):
            continue
        for line in payload:
            if isinstance(line, dict):
                text = line.get("text") or line.get("transcription") or line.get("rec_text")
                if text:
                    score = line.get("score", line.get("confidence", 1.0))
                    texts.append((str(text), float(score)))
                continue
            if isinstance(line, list) and len(line) >= 2:
                recognition = line[1]
                if isinstance(recognition, list | tuple) and recognition:
                    texts.append(
                        (
                            str(recognition[0]),
                            float(recognition[1] if len(recognition) > 1 else 1.0),
                        )
                    )
    return texts


def _remote_ocr(value: str | bytes, settings: DocumentSettings) -> Any:
    """Send one image as base64 JSON to the deployment-configured OCR endpoint."""
    if not settings.remote_ocr_configured:
        raise DocumentParseError("document_ocr_remote_unconfigured")
    payload = Path(value).read_bytes() if isinstance(value, str) else value
    request_payload = {
        "file": base64.b64encode(payload).decode("ascii"),
        "fileType": settings.ocr_file_type,
        "useDocOrientationClassify": True,
        "useDocUnwarping": False,
        "useTextlineOrientation": True,
        "returnWordBox": False,
        "visualize": False,
    }
    headers = (
        {"Authorization": f"Bearer {settings.ocr_api_key}"}
        if settings.ocr_api_key
        else {}
    )
    try:
        with httpx.Client(timeout=settings.ocr_timeout_seconds) as client:
            response = client.post(
                settings.ocr_url,
                json=request_payload,
                headers=headers,
            )
            response.raise_for_status()
            try:
                result = response.json()
            except ValueError as exc:
                raise DocumentParseError(
                    "document_ocr_remote_invalid_response"
                ) from exc
            if not isinstance(result, dict):
                raise DocumentParseError("document_ocr_remote_invalid_response")
            error_code = result.get("errorCode")
            if error_code not in (None, 0, "0"):
                safe_code = re.sub(r"[^a-zA-Z0-9_.-]", "_", str(error_code))
                raise DocumentParseError(f"document_ocr_remote_error_{safe_code}")
            if "result" not in result:
                raise DocumentParseError("document_ocr_remote_invalid_response")
            return result
    except httpx.HTTPStatusError as exc:
        raise DocumentParseError(
            f"document_ocr_remote_http_{exc.response.status_code}"
        ) from exc
    except httpx.TimeoutException as exc:
        raise DocumentParseError("document_ocr_remote_timeout") from exc
    except (httpx.HTTPError, OSError) as exc:
        raise DocumentParseError("document_ocr_remote_unavailable") from exc


def _ocr_image(value: str | bytes, settings: DocumentSettings) -> tuple[str, float]:
    if settings.ocr_mode == "remote":
        result = _remote_ocr(value, settings)
    else:
        engine = _paddle_engine()
        if hasattr(engine, "predict"):
            result = engine.predict(value)
        else:
            result = engine.ocr(value, cls=True)
    lines = [
        (text.strip(), score)
        for text, score in _texts_from_ocr_result(result)
        if text.strip()
    ]
    accepted = [text for text, score in lines if score >= 0.45]
    confidence = sum(score for _, score in lines) / len(lines) if lines else 0.0
    return "\n".join(accepted), confidence


def _parse_pdf(path: Path, settings: DocumentSettings) -> ParsedDocument:
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise DocumentParseError("document_pdf_invalid") from exc
    if document.needs_pass:
        document.close()
        raise DocumentParseError("document_encrypted_not_supported")
    if document.page_count > settings.max_logical_units:
        document.close()
        raise DocumentParseError("document_logical_unit_limit_exceeded")
    parsed = ParsedDocument(page_count=document.page_count)
    try:
        for index, page in enumerate(document):
            text = page.get_text("text", sort=True).strip()
            if len(re.sub(r"\s+", "", text)) < 50 and settings.ocr_enabled:
                scale = settings.ocr_dpi / 72
                pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                text, confidence = _ocr_image(pixmap.tobytes("png"), settings)
                parsed.ocr_pages += 1
                if confidence < 0.6:
                    parsed.quality_flags.append(f"low_ocr_confidence:page:{index + 1}")
            if text.strip():
                parsed.units.append(
                    ParsedUnit(locator=f"page:{index + 1}", text=text.strip())
                )
    finally:
        document.close()
    return parsed


def _parse_image(path: Path, settings: DocumentSettings) -> ParsedDocument:
    if not settings.ocr_enabled:
        raise DocumentParseError("document_ocr_disabled")
    text, confidence = _ocr_image(str(path), settings)
    flags = ["low_ocr_confidence:image"] if confidence < 0.6 else []
    return ParsedDocument(
        units=[ParsedUnit(locator="image:1", text=text)] if text else [],
        page_count=1,
        ocr_pages=1,
        quality_flags=flags,
    )


def _parse_docx(path: Path, settings: DocumentSettings) -> ParsedDocument:
    try:
        from docx import Document

        document = Document(path)
    except Exception as exc:
        raise DocumentParseError("document_docx_invalid") from exc
    units: list[ParsedUnit] = []
    heading: str | None = None
    block: list[str] = []

    def flush() -> None:
        if block:
            units.append(
                ParsedUnit(
                    locator=f"section:{len(units) + 1}",
                    heading=heading,
                    text="\n".join(block),
                )
            )
            block.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            flush()
            heading = text
        else:
            block.append(text)
    flush()
    for table_index, table in enumerate(document.tables, 1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells) for row in table.rows
        ]
        if any(row.strip(" |") for row in rows):
            units.append(
                ParsedUnit(
                    locator=f"table:{table_index}",
                    heading=f"Table {table_index}",
                    text="\n".join(rows),
                )
            )
    if len(units) > settings.max_logical_units:
        raise DocumentParseError("document_logical_unit_limit_exceeded")
    return ParsedDocument(units=units, page_count=len(units))


def _parse_pptx(path: Path, settings: DocumentSettings) -> ParsedDocument:
    try:
        from pptx import Presentation

        presentation = Presentation(path)
    except Exception as exc:
        raise DocumentParseError("document_pptx_invalid") from exc
    if len(presentation.slides) > settings.max_logical_units:
        raise DocumentParseError("document_logical_unit_limit_exceeded")
    units = []
    for index, slide in enumerate(presentation.slides, 1):
        lines: list[str] = []
        title: str | None = None
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip() or None
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value and value != title:
                    lines.append(value)
            if getattr(shape, "has_table", False):
                lines.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in shape.table.rows
                )
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes:\n{notes}")
        except (AttributeError, KeyError):
            pass
        text = "\n".join(lines).strip()
        if title or text:
            units.append(
                ParsedUnit(
                    locator=f"slide:{index}", heading=title, text=text or title or ""
                )
            )
    return ParsedDocument(units=units, page_count=len(presentation.slides))


def _parse_xlsx(path: Path, settings: DocumentSettings) -> ParsedDocument:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise DocumentParseError("document_xlsx_invalid") from exc
    if len(workbook.sheetnames) > settings.max_logical_units:
        workbook.close()
        raise DocumentParseError("document_logical_unit_limit_exceeded")
    units: list[ParsedUnit] = []
    try:
        for sheet in workbook.worksheets:
            batch: list[str] = []
            start_row = 1
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [str(value) if value is not None else "" for value in row]
                if any(values):
                    batch.append(" | ".join(values))
                if len(batch) >= 100:
                    units.append(
                        ParsedUnit(
                            locator=f"sheet:{sheet.title}:rows:{start_row}-{row_index}",
                            heading=sheet.title,
                            text="\n".join(batch),
                        )
                    )
                    batch = []
                    start_row = row_index + 1
                if row_index > 100_000:
                    raise DocumentParseError("document_row_limit_exceeded")
            if batch:
                units.append(
                    ParsedUnit(
                        locator=f"sheet:{sheet.title}:rows:{start_row}-{sheet.max_row}",
                        heading=sheet.title,
                        text="\n".join(batch),
                    )
                )
    finally:
        workbook.close()
    return ParsedDocument(units=units, page_count=len(workbook.sheetnames))


def _decode_text(path: Path) -> str:
    raw = path.read_bytes()
    match = from_bytes(raw).best()
    if match is None:
        raise DocumentParseError("document_text_encoding_unknown")
    return str(match)


def _parse_text(path: Path, settings: DocumentSettings) -> ParsedDocument:
    del settings
    text = _decode_text(path).strip()
    return ParsedDocument(
        units=[ParsedUnit(locator="text:1", text=text)] if text else [], page_count=1
    )


def _parse_csv(path: Path, settings: DocumentSettings) -> ParsedDocument:
    del settings
    content = _decode_text(path)
    reader = csv.reader(io.StringIO(content))
    units: list[ParsedUnit] = []
    rows: list[str] = []
    start = 1
    for index, row in enumerate(reader, 1):
        rows.append(" | ".join(cell.strip() for cell in row))
        if len(rows) >= 100:
            units.append(
                ParsedUnit(locator=f"rows:{start}-{index}", text="\n".join(rows))
            )
            rows = []
            start = index + 1
        if index > 100_000:
            raise DocumentParseError("document_row_limit_exceeded")
    if rows:
        units.append(
            ParsedUnit(
                locator=f"rows:{start}-{start + len(rows) - 1}", text="\n".join(rows)
            )
        )
    return ParsedDocument(units=units, page_count=len(units))


def parse_document(
    path: Path, media_type: str, settings: DocumentSettings
) -> ParsedDocument:
    """Parse a validated document into bounded, source-located units."""
    if media_type == "application/pdf":
        parsed = _parse_pdf(path, settings)
    elif (
        media_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        parsed = _parse_docx(path, settings)
    elif (
        media_type
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ):
        parsed = _parse_pptx(path, settings)
    elif (
        media_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        parsed = _parse_xlsx(path, settings)
    elif media_type == "text/csv":
        parsed = _parse_csv(path, settings)
    elif media_type in {"text/plain", "text/markdown"}:
        parsed = _parse_text(path, settings)
    elif media_type.startswith("image/"):
        parsed = _parse_image(path, settings)
    else:
        raise DocumentParseError("document_type_not_supported")
    if not any(unit.text.strip() for unit in parsed.units):
        raise DocumentParseError("document_no_usable_text")
    return parsed
